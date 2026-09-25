"""
to_docx.py -- Tier 3: turn extracted Page objects into a real .docx file.

NO MODEL IS INVOLVED HERE, and that is deliberate.

A .docx is not a text file. It is a zip archive of XML parts held together by
relationship ids, content-type declarations and style references that must all
resolve. Asking any language model -- small or large -- to emit that directly
is a known way to produce a file that looks plausible in a text dump and then
refuses to open. python-docx builds the real thing from a data structure, so
the output is openable by construction.

So the division of labour is:

    the model   decides WHAT is on the page   (blocks, tables, text)
    this file   decides HOW it is written out (Word styles, real table grids)

The model's job ends at the Page object. Nothing here asks it anything.

--------------------------------------------------------------------------
WHAT THIS HAS TO DEFEND AGAINST

The extraction is not guaranteed clean, and Word is less forgiving than a
terminal print. Specifically:

  - rows that disagree with the header width. Observed repeatedly: a table
    claiming 6 columns with a 7-cell row. python-docx would raise IndexError,
    so rows are padded or truncated to the grid and the damage is reported
    rather than crashing the run.

  - spanning divider rows (text in cell 1, the rest empty). These are merged
    back into a single full-width cell, which is what they were in the PDF.

  - empty headers, tables with no rows, blocks with no text. All skipped
    rather than producing an empty husk of a table.

verify_docx() then reads the finished file back and compares it against the
PDF's own text, because "it wrote a file" and "the file contains the document"
are different claims.
"""

import hashlib
from collections import Counter
from typing import List

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def apply_layout(doc, layout: dict):
    """Match the Word page to the source PDF's geometry.

    Without this, a 2-page PDF becomes a 4-page Word document -- not because
    anything was added, but because Word's defaults are roomier than the
    source at every level. Measured against one real report:

        text column   PDF 6.56 in   Word 6.00 in   (1.25 in default margins)
        body font     PDF 10 pt     Word 11 pt
        line spacing  PDF 1.0       Word 1.08
        space after   PDF ~0        Word 8 pt on EVERY paragraph

    Each is small; compounded over a page of text they add 40-50% of vertical
    space. Matching them does not guarantee identical pagination -- reflowed
    text never breaks in exactly the same places -- but it removes the
    systematic inflation, which is what turns 2 pages into 4.
    """
    section = doc.sections[0]
    section.page_width = Pt(layout["width_pt"])
    section.page_height = Pt(layout["height_pt"])
    section.left_margin = Pt(layout["left_pt"])
    section.right_margin = Pt(layout["right_pt"])
    section.top_margin = Pt(layout["top_pt"])
    section.bottom_margin = Pt(layout["bottom_pt"])

    normal = doc.styles["Normal"]
    normal.font.size = Pt(layout["body_pt"])
    pf = normal.paragraph_format
    pf.line_spacing = 1.0
    pf.space_after = Pt(2)
    pf.space_before = Pt(0)

    # Word's heading styles carry generous space_before, which on a document
    # with several consecutive headings pushes content onto another page.
    for name in ("Heading 1", "Heading 2", "Heading 3", "Title"):
        try:
            hpf = doc.styles[name].paragraph_format
        except KeyError:
            continue
        hpf.space_before = Pt(6)
        hpf.space_after = Pt(2)
        hpf.line_spacing = 1.0

from blocks import Page, TableBlock, normalize, squash

# Which Word style each block kind becomes. Keeping this as a table rather
# than a chain of ifs makes it obvious what is supported and easy to change.
HEADING_LEVELS = {"heading": 1}


def _add_table(doc, block: TableBlock, problems: List[str], where: str):
    """Write one TableBlock as a real Word table."""
    if not block.rows:
        problems.append(f"{where}: table skipped (no rows)")
        return

    # A headerless table -- a plain grid of numbers, say -- takes its width
    # from the data instead, and gets no bold first row.
    width = len(block.header) if block.has_header else len(block.rows[0])
    if width == 0:
        problems.append(f"{where}: table skipped (no columns)")
        return

    if block.has_header:
        table = doc.add_table(rows=1, cols=width)
        table.style = "Table Grid"
        for i, label in enumerate(block.header[:width]):
            cell = table.rows[0].cells[i]
            cell.text = label
            for para in cell.paragraphs:
                for run in para.runs:
                    run.bold = True
    else:
        table = doc.add_table(rows=0, cols=width)
        table.style = "Table Grid"

    for r, row in enumerate(block.rows):
        # Force the row onto the grid. A ragged row is a real extraction fault,
        # but crashing the conversion helps nobody -- record it and continue.
        if len(row) != width:
            problems.append(
                f"{where} row {r}: had {len(row)} cells, grid is {width} "
                f"-- {'padded' if len(row) < width else 'truncated'}"
            )
        cells = (list(row) + [""] * width)[:width]

        out = table.add_row().cells
        for i, value in enumerate(cells):
            out[i].text = value

        # A row with text only in the first cell is a spanning divider in the
        # source. Merge it back so the Word table looks like the PDF did.
        if cells[0] and not any(c.strip() for c in cells[1:]) and width > 1:
            merged = out[0].merge(out[width - 1])
            merged.text = cells[0]
            for para in merged.paragraphs:
                for run in para.runs:
                    run.bold = True


ALIGNMENTS = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
}


def _unescape(text: str) -> str:
    r"""Turn literal backslash-n into a real line break.

    The model sometimes emits the two characters \ and n instead of a newline,
    so a five-line address arrived as one run-on line reading
    "Dr. P.N. Cundall,\nMining Surveys Ltd.,\nHolroyd Road,...". The schema now
    asks for real breaks, but a prose instruction is advisory -- this repairs
    whatever still slips through rather than trusting it not to.
    """
    return text.replace("\\r\\n", "\n").replace("\\n", "\n").replace("\\t", "\t")


def _write_lines(para, text, italic=False, size=None):
    """Write text into a paragraph, keeping its line breaks as line breaks.

    Word treats a newline inside a run as nothing at all, so the lines of an
    address would collapse into one. Each line becomes its own run with an
    explicit break between them, which keeps them visually separate while
    staying a single paragraph -- which is what they are on the page.
    """
    lines = _unescape(text).split("\n")
    for i, line in enumerate(lines):
        if i:
            para.add_run().add_break()
        run = para.add_run(line)
        if italic:
            run.italic = True
        if size:
            run.font.size = size
    return para


def _size_scale(pages):
    """Work out what each size label should mean for THIS document.

    The model labels blocks large/normal/small, but its anchor drifts: on a
    scanned letter it marked every block except the letterhead as 'small', so
    the entire body rendered at 8.5pt.

    Prose in the field description did not fix that, so it is settled by
    measurement instead: whichever label most blocks carry IS the body size,
    by definition. That is the same rule _body_size() uses on the digital path,
    where the dominant font size on the page defines the body. The other labels
    scale relative to whatever that turns out to be.
    """
    base = {"small": 0.8, "normal": 1.0, "large": 1.6}

    # Weighted by CHARACTERS, not by block count. The body of a document is
    # whatever carries the most text, which is the same rule the digital path
    # uses to find the body font size. Counting blocks instead let a page with
    # one long paragraph and three short headings decide that headings were
    # the body -- on one scan that made 'large' the anchor and collapsed the
    # actual body text to 0.5x, rendering it at 5.5pt.
    counts = Counter()
    for page in pages:
        for b in page.blocks:
            if isinstance(b, TableBlock):
                continue
            counts[getattr(b, "size", "normal")] += len(b.text)

    if not counts:
        return base

    anchor = base[counts.most_common(1)[0][0]]
    return {k: v / anchor for k, v in base.items()}


def _apply_look(para, block, layout, is_heading=False, scale=None):
    """Apply the block's observed appearance: alignment, size, weight.

    These come from the model looking at the page, because on a scan there is
    no other source -- no text layer means no font sizes and no coordinates.
    Sizes are relative to the document's body size so the result scales with
    whatever the source used.
    """
    para.alignment = ALIGNMENTS.get(getattr(block, "align", "left"),
                                    WD_ALIGN_PARAGRAPH.LEFT)

    body = (layout or {}).get("body_pt", 11)
    scale = scale or {"large": 1.6, "normal": 1.0, "small": 0.8}
    factor = scale.get(getattr(block, "size", "normal"), 1.0)

    # A floor and ceiling, because the labels are the model's judgement and it
    # can be badly skewed. Whatever it reports, body text must stay readable
    # and a heading must not become a billboard: one scan produced 5.5pt
    # paragraphs before this existed. 8pt is small print; nothing legitimate
    # in a document needs to be smaller.
    point = max(8.0, min(body * factor, body * 2.2))

    for run in para.runs:
        if getattr(block, "bold", False):
            run.bold = True
        # A heading style already carries its own size; only override when the
        # model saw something other than ordinary body text.
        if factor != 1.0 or not is_heading:
            if run.font.size is None or factor != 1.0:
                run.font.size = Pt(round(point, 1))


def _add_images(doc, images, layout, problems, where):
    """Insert extracted pictures at their displayed size.

    Width is taken from how big the image was ON THE PAGE, not its pixel
    dimensions, so a 1352x969 logo shown at 6.13in comes out 6.13in rather
    than 18in. It is then clamped to the text column, because a picture wider
    than the margins silently overflows the page in Word.
    """
    max_in = 6.5
    if layout:
        max_in = (layout["width_pt"] - layout["left_pt"]
                  - layout["right_pt"]) / 72.0

    for im in images:
        width = min(im["width_in"], max_in)
        try:
            doc.add_picture(im["path"], width=Inches(width))
        except Exception as exc:
            problems.append(f"{where}: could not insert {im['path']}: {exc}")


def build_docx(pages: List[Page], path: str, title: str = None,
               layout: dict = None, image_sets: List[list] = None):
    """Write extracted pages to a .docx. Returns a list of problems found.

    pages is one Page per PDF page, in order. Each gets a page break after it,
    so the Word document keeps the source's pagination.

    layout, when given, matches the Word page geometry and text size to the
    source PDF -- see apply_layout(). Without it Word's roomier defaults
    inflate the document onto extra pages.

    image_sets, when given, is one list of extracted images per page, each
    carrying the fraction of that page's text sitting above it. The model's
    blocks have no coordinates, so that fraction is what places a picture in
    reading order: "65% of the text is above this logo" becomes "insert it
    after 65% of the blocks". A page with no images contributes an empty list
    and nothing about its output changes.
    """
    problems: List[str] = []
    doc = Document()

    # What "normal" means is decided by the document, not by the label.
    scale = _size_scale(pages)

    if layout:
        apply_layout(doc, layout)

    # The filename heading is skipped when matching the source layout: the PDF
    # has no such line, so adding one guarantees the first page differs.
    if title and not layout:
        doc.add_heading(title, level=0)

    for pno, page in enumerate(pages, start=1):
        images = (image_sets[pno - 1] if image_sets
                  and pno - 1 < len(image_sets) else [])

        # Turn each image's "fraction of text above me" into a block index to
        # sit after. Sorted so several images on one page stay in page order,
        # and popped from the front as the blocks are written out.
        n = len(page.blocks)
        # key=... is load-bearing, not style: with an empty page (n == 0 --
        # exactly the separate scan document under --scan-mode both) every
        # image's position collapses to 0, so every tuple ties on its first
        # element. Without a key, Python falls back to comparing the second
        # element (the image dict) to break the tie, and dicts have no
        # ordering -- a real crash, reproduced with 2 images on one blank
        # page. The key restricts comparison to the position only; Python's
        # sort is stable, so tied images simply keep their original order.
        pending = sorted(
            ((min(n, round(im["frac_above"] * n)), im) for im in images),
            key=lambda t: t[0],
        )

        # A table is a different kind of element than a paragraph in Word's
        # own format, and does not inherit the Normal style's space_after the
        # way one paragraph inherits it from the paragraph before it. Left
        # alone, a table sits flush against the text on both sides -- true on
        # every page, not something specific to any one document. Tracking
        # the most recently written paragraph, and whether a table was just
        # written, is what lets the gap be added explicitly on both sides.
        last_para = None
        after_table = False

        for bno, block in enumerate(page.blocks):
            where = f"page {pno} block[{bno}]"

            while pending and pending[0][0] <= bno:
                _add_images(doc, [pending.pop(0)[1]], layout, problems, where)

            if isinstance(block, TableBlock):
                if last_para is not None:
                    last_para.paragraph_format.space_after = Pt(6)
                _add_table(doc, block, problems, where)
                after_table = True
                continue

            text = block.text.strip()
            if not text:
                problems.append(f"{where}: empty {block.kind}, skipped")
                continue

            if block.kind in HEADING_LEVELS:
                para = doc.add_heading("", level=HEADING_LEVELS[block.kind])
                _write_lines(para, text)
                _apply_look(para, block, layout, is_heading=True, scale=scale)

            elif block.kind == "list":
                # The model returns a list as one text block; split it back
                # into bullets on line breaks and common bullet characters.
                items = [
                    line.strip(" •◦▪-–\t")
                    for line in _unescape(text).splitlines()
                    if line.strip(" •◦▪-–\t")
                ]
                para = None
                for item in items or [text]:
                    para = doc.add_paragraph(item, style="List Bullet")

            elif block.kind == "caption":
                para = doc.add_paragraph()
                _write_lines(para, text, italic=True, size=Pt(9))
                _apply_look(para, block, layout, scale=scale)

            else:  # paragraph, image, anything else
                para = doc.add_paragraph()
                _write_lines(para, text)
                _apply_look(para, block, layout, scale=scale)

            if after_table and para is not None:
                para.paragraph_format.space_before = Pt(6)
            after_table = False
            if para is not None:
                last_para = para

        # Anything whose fraction landed past the last block goes at the end.
        if pending:
            _add_images(doc, [im for _, im in pending], layout, problems,
                        f"page {pno} end")

        if pno < len(pages):
            doc.add_page_break()

    doc.save(path)
    return problems


def verify_images(path: str, image_sets):
    """Check 5: did the pictures actually survive into the Word file?

    Nothing checked this before. An image could fail to insert, insert twice,
    or land in the wrong place, and every other check would still pass --
    checks 1-4 are all about text, and verify_docx() counts words.

    Three things, all read back from the SAVED file rather than from what we
    meant to write:

      COUNT     one inline shape per extracted image. Catches a silent
                insertion failure or a duplicate.
      BYTES     sha1 of the image stored in the .docx against the sha1 of the
                bytes taken out of the PDF. Proves it is the same picture,
                not a re-encoded or mismatched one.
      POSITION  the coordinate check. The PDF says what text sits immediately
                above and below the image; this confirms the same text sits
                either side of it in the document. frac_above is a heuristic,
                and until now its result was never verified -- "roughly the
                right place" becomes pass or fail.

    Returns a list of problems; empty means all three passed.
    """
    expected = [im for page in (image_sets or []) for im in page]
    doc = Document(path)
    shapes = doc.inline_shapes

    problems = []

    if len(shapes) != len(expected):
        problems.append(
            f"COUNT: {len(expected)} image(s) extracted from the PDF but "
            f"{len(shapes)} in the .docx"
        )

    # --- bytes -----------------------------------------------------------
    for i, (shape, im) in enumerate(zip(shapes, expected)):
        try:
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            blob = doc.part.related_parts[rid].blob
        except Exception as exc:
            problems.append(f"BYTES: image {i} unreadable in the .docx: {exc}")
            continue
        got = hashlib.sha1(blob).hexdigest()
        if got != im.get("sha1"):
            problems.append(
                f"BYTES: image {i} differs -- PDF sha1 {im.get('sha1', '?')[:12]}, "
                f".docx sha1 {got[:12]}"
            )

    # --- position --------------------------------------------------------
    # Walk the body in order, recording text and where the pictures fall.
    sequence = []
    for child in doc.element.body.iterchildren():
        if child.tag == f"{W}p":
            para = Paragraph(child, doc)
            if para._p.findall(f".//{W}drawing"):
                sequence.append(("image", None))
            elif para.text.strip():
                sequence.append(("text", normalize(para.text)))
        elif child.tag == f"{W}tbl":
            table = Table(child, doc)
            cells = " ".join(c.text for r in table.rows for c in r.cells)
            sequence.append(("text", normalize(cells)))

    positions = [i for i, (kind, _) in enumerate(sequence) if kind == "image"]

    for i, (pos, im) in enumerate(zip(positions, expected)):
        want_before = im.get("text_before")
        want_after = im.get("text_after")

        if want_before:
            prev = next((t for k, t in reversed(sequence[:pos]) if k == "text"), "")
            if normalize(want_before)[:40] not in prev:
                problems.append(
                    f"POSITION: image {i} should follow "
                    f"{want_before[:40]!r} but follows {prev[:40]!r}"
                )
        if want_after:
            nxt = next((t for k, t in sequence[pos + 1:] if k == "text"), "")
            if normalize(want_after)[:40] not in nxt:
                problems.append(
                    f"POSITION: image {i} should precede "
                    f"{want_after[:40]!r} but precedes {nxt[:40]!r}"
                )

    return problems


def verify_docx(path: str, source_text: str):
    """Read the written .docx back and check it against the PDF's own text.

    "A file was produced" and "the file contains the document" are different
    claims. This checks the second one, by reading every paragraph and every
    table cell out of the saved file and measuring how much of the PDF's text
    is accounted for.

    Returns (coverage, missing_words). Coverage is the fraction of the PDF's
    words (longer than 3 characters) that appear somewhere in the document.
    """
    doc = Document(path)

    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    written = normalize(" ".join(parts))
    written_squashed = squash(" ".join(parts))

    words = [w for w in normalize(source_text).split() if len(w) > 3]
    if not words:
        return 1.0, []

    # squash() catches words that differ only in spacing, such as
    # {"model","method"} against {"model", "method"}.
    missing = sorted({
        w for w in words
        if w not in written and squash(w) not in written_squashed
    })
    coverage = 1 - len(missing) / len(set(words))
    return coverage, missing
