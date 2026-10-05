"""Regression tests that never call the model. Each one guards a fix from FIX_LOG.md."""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

import _load
from _load import AGREEMENT, AGREEMENT_SCANNED, REPO

import pymupdf

from blocks import LayoutAnalysis, Page, TableBlock, TextBlock, normalize

HAVE_TESSERACT = bool(shutil.which("tesseract"))
HAVE_SOFFICE = bool(shutil.which("soffice"))


def text(kind, t):
    return TextBlock(kind=kind, text=t, align="left", size="normal", bold=False)


def table(header, rows):
    return TableBlock(kind="table", has_header=True, n_data_rows=len(rows), n_cols=len(header), header=header, rows=rows)


def page_of(blocks):
    counts = [b.n_cols for b in blocks if isinstance(b, TableBlock)]
    return Page(analysis=LayoutAnalysis(n_blocks=len(blocks), table_column_counts=counts), blocks=blocks)


# --------------------------------------------------------------------------- fix #2
class ContentCheckLine(unittest.TestCase):
    """KNOWN_ISSUES #2: a scan must not report "100%" against an empty text layer."""

    @classmethod
    def setUpClass(cls):
        cls.f = staticmethod(_load.load_script_defs("test_pdf.py").content_check_line)

    def test_all_scanned_says_not_run(self):
        self.assertIn("not run", self.f(1.0, "docx", ["", " "]))

    def test_digital_keeps_the_percentage(self):
        self.assertIn("98%", self.f(0.98, "docx", ["real text", "more"]))

    def test_mixed_names_the_scanned_pages(self):
        line = self.f(0.97, "pptx", ["real text", ""])
        self.assertIn("97%", line)
        self.assertIn("1 scanned page", line)


# --------------------------------------------------------------------------- fix #5
class PowerPointTitles(unittest.TestCase):
    """KNOWN_ISSUES #5: a slide title is a short line; merged heading blocks are split."""

    def kinds(self, block):
        from to_pptx import _split_headings
        return [(b.kind, b.text) for b in _split_headings([block])]

    def test_title_and_subtitle_split(self):
        out = self.kinds(text("heading", "Services Agreement\nSample document for translation testing"))
        self.assertEqual(out[0], ("heading", "Services Agreement"))
        self.assertEqual(out[1][0], "paragraph")

    def test_clause_heading_inside_a_paragraph_splits(self):
        out = self.kinds(text("paragraph", "3. Fees\nThe Buyer shall pay the fees below. All amounts are in rupees."))
        self.assertEqual([k for k, _ in out], ["heading", "paragraph"])

    def test_things_that_must_stay_whole(self):
        for block in (text("heading", "Annual Report\n2026"),
                      text("list", "1. Apples\n2. Pears and a much longer second line here"),
                      text("paragraph", "The Buyer shall pay the fees\nbelow within thirty days of the invoice."),
                      text("paragraph", "2. Services are provided as described.\nThe rest follows and is longer than that."),
                      text("paragraph", "Page 1 of 2")):
            self.assertEqual(len(self.kinds(block)), 1, block.text)

    def test_deck_titles_are_short_lines(self):
        from pptx import Presentation
        from to_pptx import build_pptx
        pages = [page_of([text("heading", "Services Agreement\nSample document\nACME"), text("paragraph", "Body.")]),
                 page_of([text("heading", "4. Payment\n" + "The Buyer shall pay each invoice within 45 days. " * 3)])]
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "deck.pptx")
            self.assertEqual(build_pptx(pages, path), [])
            titles = [s.shapes.title.text for s in Presentation(path).slides]
        self.assertEqual(titles, ["Services Agreement", "4. Payment"])


# --------------------------------------------------------------------------- fix #1
@unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
class CropAndReread(unittest.TestCase):
    """KNOWN_ISSUES #1: the model stopped early; the missing region is cropped and read again."""

    @classmethod
    def setUpClass(cls):
        import scan_fixtures
        import crop_reread
        from blocks import check_coverage
        import ocr
        cls.cr, cls.check_coverage, cls.ocr = crop_reread, staticmethod(check_coverage), ocr
        cls.tmp = tempfile.mkdtemp()
        cls.exp = {}
        cls.pdf = {}
        for fn in (scan_fixtures.table_low, scan_fixtures.table_high, scan_fixtures.two_tables,
                   scan_fixtures.no_table, scan_fixtures.wide_table):
            name = f"scan_{fn.__name__}.pdf"
            path = os.path.join(cls.tmp, name)
            cls.exp[name] = fn(path)
            cls.pdf[name] = path

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def image(self, name):
        path = os.path.join(self.tmp, name[:-4] + ".png")
        pymupdf.open(self.pdf[name])[0].get_pixmap(dpi=200).save(path)
        return path

    def missing(self, name, page):
        return self.check_coverage(page, self.ocr.ocr_page(pymupdf.open(self.pdf[name])[0]))[3]

    def reader(self, blocks):
        calls = []

        def read(path, model, base_url, include_look):
            calls.append(path)
            return page_of(blocks)
        read.calls = calls
        return read

    def broken_page(self):
        e = self.exp["scan_table_low.pdf"]
        return page_of([text("paragraph", l) for l in e["lines"] if l != "Page 1 of 2"]), e

    def test_table_and_footer_come_back_in_order(self):
        page, e = self.broken_page()
        rd = self.reader([table(**{"header": e["tables"][0]["header"], "rows": e["tables"][0]["rows"]}),
                          text("paragraph", "Page 1 of 2")])
        fixed, notes = self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=rd,
                                       missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertEqual(len(rd.calls), 1)
        self.assertEqual([b.kind for b in fixed.blocks][-2:], ["table", "paragraph"])
        self.assertEqual(len(fixed.blocks), len(page.blocks) + 2)
        self.assertTrue(notes)

    def test_a_second_round_picks_up_a_leftover_line(self):
        """The model returned the table but wrote the footer as an empty paragraph; the footer gets its own crop."""
        page, e = self.broken_page()
        answers = [[text("paragraph", ""), table(e["tables"][0]["header"], e["tables"][0]["rows"])],
                   [text("paragraph", "Page 1 of 2")]]
        calls = []

        def read(path, model, base_url, include_look):
            calls.append(path)
            return page_of(answers[len(calls) - 1])
        fixed, notes = self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=read,
                                       missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertEqual(len(calls), 2, "one call per round")
        self.assertEqual([b.kind for b in fixed.blocks][-2:], ["table", "paragraph"])
        self.assertEqual(fixed.blocks[-1].text, "Page 1 of 2")
        self.assertTrue(all(getattr(b, "text", "x").strip() for b in fixed.blocks), "no empty block was added")
        self.assertTrue(any(n.startswith("round 2") for n in notes))

    def test_it_stops_after_two_rounds(self):
        page, _ = self.broken_page()
        rd = self.reader([text("paragraph", "Nothing useful here at all, unrelated to this page.")])
        self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=rd,
                        missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertLessEqual(len(rd.calls), 2)

    def test_complete_pages_are_never_touched(self):
        for name, e in self.exp.items():
            full = page_of([text("paragraph", l) for l in e["lines"]]
                           + [table(t["header"], t["rows"]) for t in e["tables"]])
            ml = self.missing(name, full)
            self.assertEqual(ml, [], f"{name}: the coverage check should flag nothing")

    def test_a_useless_reread_is_rejected(self):
        page, _ = self.broken_page()
        rd = self.reader([text("paragraph", "A paragraph about something else that is not on this page at all.")])
        same, _ = self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=rd,
                                  missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertEqual(len(same.blocks), len(page.blocks))

    def test_repeated_text_is_not_added_twice(self):
        page, e = self.broken_page()
        rd = self.reader([text("heading", "3. Charges"), table(e["tables"][0]["header"], e["tables"][0]["rows"]),
                          text("paragraph", "Page 1 of 2")])
        fixed, _ = self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=rd,
                                   missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertEqual([normalize(getattr(b, "text", "")) for b in fixed.blocks].count("3. charges"), 1)

    def test_a_model_error_leaves_the_page_alone(self):
        page, _ = self.broken_page()

        def boom(path, model, base_url, include_look):
            raise RuntimeError("model server down")
        same, notes = self.cr.recover(page, self.image("scan_table_low.pdf"), "m", "u", True, reader=boom,
                                      missing_lines=self.missing("scan_table_low.pdf", page))
        self.assertEqual(len(same.blocks), len(page.blocks))
        self.assertIn("could not read", notes[0])


# --------------------------------------------------------------------------- fix #6
class WordTable(unittest.TestCase):
    """KNOWN_ISSUES #6: the table is indented by the cell padding and rows keep the PDF's height."""

    def build(self):
        import pdf_look
        from to_docx import build_docx
        H = _load.load_script_defs("test_pdf.py")
        doc = pymupdf.open(AGREEMENT)
        p1 = page_of([text("heading", "Services Agreement"), text("paragraph", "3. Fees\nThe Buyer shall pay the fees below."),
                      table(["Item", "Amount (INR)"], [["Setup and configuration", "25,000.00"],
                                                       ["Monthly service fee", "15,000.00"], ["Training (2 days)", "8,000.00"]])])
        images = [[], []]
        layout = H.measure_layout(doc[0])
        pages, lay, _ = pdf_look.restyle_document([p1, page_of([text("paragraph", "Page 2 of 2")])], [doc[0], doc[1]],
                                                  [True, True], images, layout)
        self.tmp = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, self.tmp, True)
        path = os.path.join(self.tmp, "t.docx")
        build_docx(pages, path, layout=lay, image_sets=images)
        return path, doc

    def test_indent_and_row_height_are_written(self):
        import docx
        from docx.oxml.ns import qn
        path, _ = self.build()
        tbl = docx.Document(path).tables[0]
        ind = tbl._tbl.tblPr.find(qn("w:tblInd"))
        self.assertIsNotNone(ind)
        self.assertEqual(ind.get(qn("w:w")), "108")
        names = [c.tag.split("}")[1] for c in tbl._tbl.tblPr]
        self.assertLess(names.index("tblInd"), names.index("tblLook"), "the schema's element order must hold")

    @unittest.skipUnless(HAVE_SOFFICE, "needs LibreOffice")
    def test_table_lines_up_with_the_pdf(self):
        path, doc = self.build()
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", self.tmp, path],
                       capture_output=True, check=True)
        out = pymupdf.open(os.path.join(self.tmp, "t.pdf"))
        a, b = doc[0].find_tables().tables[0], out[0].find_tables().tables[0]
        self.assertAlmostEqual(b.bbox[0], a.bbox[0], delta=1.0, msg="left edge within 1 pt of the PDF's")
        self.assertAlmostEqual(b.bbox[2] - b.bbox[0], a.bbox[2] - a.bbox[0], delta=1.5)


# --------------------------------------------------------------------------- fix #14
@unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
class RepeatedLineFlag(unittest.TestCase):
    """KNOWN_ISSUES #14: a line written twice but read once by OCR is pointed out; nothing is removed."""

    def verify(self, body):
        import ocr
        M = _load.load_script_defs("pdf_to_txt.py")
        doc = pymupdf.open(AGREEMENT_SCANNED)
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "t.txt")
            open(path, "w").write(body)
            results, _ = M.verify_txt(path, doc, [(1, "scanned")], {1: ocr.ocr_page(doc[0])})
        return results[0]

    def test_a_repeat_is_flagged_and_a_clean_page_is_not(self):
        clean = ("----- Page 1 -----\n\nServices Agreement\nSample document for translation testing\nACME\n\n"
                 "1. Parties\nThis services agreement is made between ACME Test Ltd (the \"Supplier\") and Example Buyer "
                 "Pvt Ltd (the \"Buyer\"). All names in this document are invented for testing.\n\nPage 1 of 2\n")
        self.assertEqual(self.verify(clean).get("repeats"), [])
        flagged = self.verify(clean + "ACME\n")
        self.assertEqual([r[0] for r in flagged["repeats"]], ["acme"])
        self.assertEqual(flagged["problems"] == [], True, "advisory only: it must not turn PASS into FAIL")


# --------------------------------------------------------------------------- API
class ApiDocs(unittest.TestCase):
    """Swagger shapes and the test page's cache header; the endpoints still return what they did."""

    @classmethod
    def setUpClass(cls):
        cls._tmp = tempfile.mkdtemp()
        os.environ["IDOX_JOBS_DIR"] = cls._tmp
        sys.path.insert(0, REPO)
        from fastapi.testclient import TestClient
        from api import idox_api
        cls.client = TestClient(idox_api.app)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls._tmp, ignore_errors=True)

    def test_response_shapes_are_documented(self):
        spec = self.client.get("/openapi.json").json()
        for name in ("Health", "ModelServer", "DeleteResult", "ErrorDetail", "Job"):
            self.assertIn(name, spec["components"]["schemas"])
        self.assertIn("404", spec["paths"]["/jobs/{jid}"]["get"]["responses"])
        self.assertIn("404", spec["paths"]["/jobs/{jid}"]["delete"]["responses"])
        self.assertIn("application/octet-stream", spec["paths"]["/jobs/{jid}/files/{name}"]["get"]["responses"]["200"]["content"])

    def test_every_conversion_still_has_its_form(self):  # 17 earlier ones plus JPG to PDF (searchable picture)
        spec = self.client.get("/openapi.json").json()
        posts = [p for p, o in spec["paths"].items() if "post" in o and p.startswith("/convert/")]
        self.assertEqual(len(posts), 18)
        self.assertIn("/convert/jpg-to-pdf-searchable", posts)

    def test_test_page_is_not_cached(self):
        r = self.client.get("/ui")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.headers.get("cache-control"), "no-cache")

    def test_health_keeps_its_shape(self):
        d = self.client.get("/health").json()
        self.assertEqual(sorted(d), ["jobs", "model_server", "tools"])
        self.assertIn("ok", d["model_server"])

    def test_unknown_job_is_a_404_with_a_detail(self):
        for r in (self.client.get("/jobs/nope"), self.client.delete("/jobs/nope"), self.client.get("/jobs/nope/log")):
            self.assertEqual(r.status_code, 404)
            self.assertIn("detail", r.json())


# --------------------------------------------------------------------------- converters that need no model
class ImageRoutes(unittest.TestCase):
    """PDF to JPG / PNG / TIFF run end to end without the model."""

    def run_script(self, script, *extra):
        with tempfile.TemporaryDirectory() as d:
            r = subprocess.run([sys.executable, os.path.join(REPO, script), AGREEMENT, "--outdir", d, *extra],
                               capture_output=True, text=True, timeout=300)
            top = os.listdir(d)
            inner = os.path.join(d, top[0]) if top else d
            files = sorted(os.listdir(inner)) if os.path.isdir(inner) else sorted(top)   # TIFF is one file
            return r, files

    def test_png_exits_cleanly_and_writes_both_pages(self):
        r, files = self.run_script("pdf_to_png.py", "--no-ocr")
        self.assertEqual(r.returncode, 0, r.stdout[-400:] + r.stderr[-400:])
        self.assertEqual(len(files), 2)

    def test_tiff_exits_cleanly_and_writes_one_file(self):
        r, files = self.run_script("pdf_to_tiff.py", "--no-ocr")
        self.assertEqual(r.returncode, 0, r.stdout[-400:] + r.stderr[-400:])
        self.assertEqual([f for f in files if f.endswith(".tiff")], ["agreement.tiff"])

    def test_jpg_exits_cleanly_and_writes_both_pages(self):
        r, files = self.run_script("pdf_to_jpg.py", "--no-ocr")
        self.assertEqual(r.returncode, 0, r.stdout[-400:] + r.stderr[-400:])
        self.assertEqual(len(files), 2)


@unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
class OcrIsAdvisory(unittest.TestCase):
    """Review item: an OCR shortfall on correctly drawn pixels is a warning, not a failure."""

    def check(self, script):
        M = _load.load_script_defs(script)
        doc = pymupdf.open(AGREEMENT)
        page = doc[0]
        ext = "jpg" if "jpg" in script else "png"
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, f"p.{ext}")
            pix = page.get_pixmap(dpi=150, alpha=False)
            pix.save(path)
            real = M.ocr.ocr_image
            M.ocr.ocr_image = lambda _p: "unreadable"          # Tesseract "reads" almost nothing
            try:
                if ext == "jpg":
                    problems, info = M.check_page(page, path, pix.width, pix.height, 150, True)
                else:
                    problems, info = M.check_page(page, path, 150, True)
            finally:
                M.ocr.ocr_image = real
        return problems, info

    def test_jpg_and_png_warn_but_do_not_fail(self):
        for script in ("pdf_to_jpg.py", "pdf_to_png.py"):
            problems, info = self.check(script)
            self.assertEqual(problems, [], f"{script}: exact checks passed, so no problem")
            self.assertIn("OCR read only", info.get("ocr_warning", ""), script)


class ScriptsAreImportable(unittest.TestCase):
    """Review item: the four smaller From-PDF scripts do nothing on import and have a main() to call."""

    def test_import_has_no_side_effects(self):
        code = ("import pdf_to_txt, pdf_to_jpg, pdf_to_png, pdf_to_tiff; "
                "print(all(callable(m.main) for m in (pdf_to_txt, pdf_to_jpg, pdf_to_png, pdf_to_tiff)))")
        r = subprocess.run([sys.executable, "-W", "ignore", "-c", code], capture_output=True, text=True, cwd=REPO, timeout=120)
        self.assertEqual(r.returncode, 0, r.stderr[-400:])
        self.assertEqual(r.stdout.strip().splitlines()[-1], "True")


class PdfToTextWithoutModel(unittest.TestCase):
    """With the model unreachable, a digital page still comes out as the PDF's own text; nothing crashes."""

    def test_falls_back_to_the_text_layer(self):
        with tempfile.TemporaryDirectory() as d:
            r = subprocess.run([sys.executable, "-W", "ignore", os.path.join(REPO, "pdf_to_txt.py"), AGREEMENT,
                                "--outdir", d, "--base-url", "http://127.0.0.1:9"],
                               capture_output=True, text=True, timeout=300)
            self.assertNotIn("Traceback", r.stderr)
            self.assertEqual(r.returncode, 0, r.stdout[-400:])
            out = open(os.path.join(d, "agreement.txt"), encoding="utf-8").read()
        self.assertIn("Services Agreement", out)
        self.assertIn("27ABCPE1234F1ZB", out)


class FalseReviewChecks(unittest.TestCase):
    """KNOWN_ISSUES #3: correct output must not be sent to review; real losses must still be."""

    def tbl(self, header, rows, n=None, has_header=True):
        return TableBlock(kind="table", has_header=has_header, n_data_rows=len(rows) if n is None else n,
                          n_cols=len(header), header=header, rows=rows)

    # -- the header row counted as a data row -------------------------------------------------------
    def test_header_miscount_is_only_a_note(self):
        from blocks import check_structure
        t = self.tbl(["Item", "Amount"], [["A", "1"], ["B", "2"], ["C", "3"]], n=4)       # model said 4, wrote 3
        problems, notes = check_structure(page_of([t]))
        self.assertEqual(problems, [])
        self.assertTrue(any("counted the header row" in n for n in notes))

    def test_other_miscounts_are_still_problems(self):
        from blocks import check_structure
        for n, has_header in ((5, True), (2, True), (4, False)):                         # +2, one fewer, no header
            t = self.tbl(["Item", "Amount"] if has_header else [], [["A", "1"], ["B", "2"], ["C", "3"]], n=n,
                         has_header=has_header)
            if not has_header:
                t = TableBlock(kind="table", has_header=False, n_data_rows=n, n_cols=2, header=[],
                               rows=[["A", "1"], ["B", "2"], ["C", "3"]])
            problems, _ = check_structure(page_of([t]))
            self.assertTrue(any("n_data_rows" in p for p in problems), (n, has_header))

    # -- OCR drops a one-digit cell from a table row ------------------------------------------------
    def test_row_tolerant_coverage(self):
        from blocks import check_coverage
        page = page_of([self.tbl(["Item", "Qty", "Amount"], [["Paper A4", "10", "450.00"], ["Toner", "2", "3,200.00"]])])
        ocr_text = "Item Qty Amount\nPaper A4 450.00\nToner 3,200.00\n"                  # Tesseract lost the quantities
        strict = check_coverage(page, ocr_text)[3]
        tolerant = check_coverage(page, ocr_text, row_tolerant=True)[3]
        self.assertTrue(strict, "the plain check still flags those lines (its behaviour is unchanged)")
        self.assertEqual(tolerant, [], "...but a row read with a digit missing is not a lost row")

    def test_a_really_missing_row_is_still_missing(self):
        from blocks import check_coverage
        page = page_of([self.tbl(["Item", "Qty", "Amount"], [["Paper A4", "10", "450.00"]])])
        ocr_text = "Item Qty Amount\nPaper A4 450.00\nStapler heavy duty 275.00\n"
        missing = check_coverage(page, ocr_text, row_tolerant=True)[3]
        self.assertEqual(missing, ["Stapler heavy duty 275.00"])

    # -- table cells that OCR cannot confirm --------------------------------------------------------
    def test_table_grounding_problems_are_split_off(self):
        from blocks import split_table_problems
        hard, soft = split_table_problems(["block[0] kind=paragraph: text not found in PDF (10% of words matched)",
                                           "block[1] kind=table: 7 cell(s) not in PDF text: '10', '2'"])
        self.assertEqual(len(hard), 1)
        self.assertEqual(len(soft), 1)

    # -- the API turns the log into a state ---------------------------------------------------------
    def api(self):
        os.environ.setdefault("IDOX_JOBS_DIR", tempfile.mkdtemp())
        from api import idox_api
        return idox_api

    LOG_HEAD = ("PAGE 1\nTier 0 -- text layer: 0 chars, tables detected: 0\n\n  --- checks ---\n"
                "  1. self-consistency:       PASS\n  2. text-layer grounding:   PASS   (8/8 strings verified)\n"
                "  4. content coverage:       FAIL   (79% of the PDF's text was extracted)\n")

    def test_a_fully_repaired_page_is_not_sent_to_review(self):
        api = self.api()
        log = self.LOG_HEAD + ("       + RESTORED by re-reading the page image: ...\n"
                               "  4. content coverage:       PASS   (100% of the PDF's text was extracted, after re-reading)\n")
        self.assertEqual(api.check_problems(api.page_checks(log)), [])

    def test_a_repair_that_did_not_finish_still_goes_to_review(self):
        api = self.api()
        log = self.LOG_HEAD + "  4. content coverage:       FAIL   (98% of the PDF's text was extracted, after re-reading)\n"
        problems = api.check_problems(api.page_checks(log))
        self.assertEqual(len(problems), 1)
        self.assertIn("coverage", problems[0])

    def test_no_repair_no_change(self):
        api = self.api()
        problems = api.check_problems(api.page_checks(self.LOG_HEAD))
        self.assertEqual(len(problems), 1)

    def test_an_advisory_table_grounding_result_is_not_a_problem(self):
        api = self.api()
        log = ("  1. self-consistency:       PASS\n  2. OCR grounding:          differs   (6/13 strings verified; advisory)\n"
               "  4. OCR coverage:           PASS   (100% of what OCR read was extracted)\n")
        self.assertEqual(api.check_problems(api.page_checks(log)), [])


def make_tiff(path, size, dpi=None, words=("PAGE ONE SAMPLE 2026", "PAGE TWO SAMPLE 4417")):
    """A made-up two-frame TIFF with large black text on white; `dpi=None` stores no resolution."""
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.load_default(size=max(24, size[1] // 6))
    frames = []
    for w in words:
        im = Image.new("RGB", size, "white")
        ImageDraw.Draw(im).text((size[0] // 20, size[1] // 3), w, fill="black", font=font)
        frames.append(im)
    kw = {"dpi": (dpi, dpi)} if dpi else {}
    frames[0].save(path, save_all=True, append_images=frames[1:], **kw)


class TiffPageSize(unittest.TestCase):
    """KNOWN_ISSUES #13: a TIFF with no stored resolution is fitted to A4, not shrunk to an assumed 300 dpi."""

    def convert(self, size, dpi=None, text=False):
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        src = os.path.join(d, "in.tiff")
        make_tiff(src, size, dpi)
        args = [sys.executable, "-W", "ignore", os.path.join(REPO, "tiff_to_pdf.py"), src, "--outdir", d]
        if not text:
            args.append("--no-text")
        r = subprocess.run(args, capture_output=True, text=True, timeout=300)
        self.assertEqual(r.returncode, 0, r.stdout[-500:] + r.stderr[-300:])
        return pymupdf.open(os.path.join(d, "in.pdf")), r.stdout

    def test_no_resolution_landscape_fits_a4(self):
        doc, out = self.convert((1000, 400))
        self.assertEqual(doc.page_count, 2)
        w, h = doc[0].rect.width, doc[0].rect.height
        self.assertAlmostEqual(w, 842, delta=1, msg="the long side fills A4's long side")
        self.assertAlmostEqual(w / h, 1000 / 400, delta=0.01, msg="never stretched")
        self.assertIn("fitted to A4", out)

    def test_no_resolution_portrait_fits_a4(self):
        doc, _ = self.convert((1000, 1400))
        w, h = doc[0].rect.width, doc[0].rect.height
        self.assertLessEqual(w, 595.5)
        self.assertLessEqual(h, 842.5)
        self.assertAlmostEqual(w / h, 1000 / 1400, delta=0.01)

    def test_an_a4_sized_scan_without_resolution_is_unchanged(self):
        doc, _ = self.convert((2480, 3508))
        self.assertAlmostEqual(doc[0].rect.width, 595, delta=1)
        self.assertAlmostEqual(doc[0].rect.height, 842, delta=1)

    def test_a_stored_resolution_is_still_respected(self):
        doc, out = self.convert((1000, 400), dpi=300)
        self.assertAlmostEqual(doc[0].rect.width, 240, delta=0.6)
        self.assertAlmostEqual(doc[0].rect.height, 96, delta=0.6)
        self.assertNotIn("fitted to A4", out)

    @unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
    def test_the_searchable_text_still_lands_on_the_page(self):
        doc, _ = self.convert((1000, 400), text=True)
        hits = doc[1].search_for("TWO")
        self.assertTrue(hits, "the hidden text layer must still contain the word")
        self.assertTrue(all(doc[1].rect.contains(h) for h in hits), "...and sit inside the page")


class JpgSearchablePicture(unittest.TestCase):
    """KNOWN_ISSUES #11: JPG to PDF that keeps the picture, fitted to A4, with a hidden OCR text layer. No model."""

    def jpg(self, density=None):
        from PIL import Image, ImageDraw, ImageFont
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        im = Image.new("RGB", (1000, 700), "white")
        ImageDraw.Draw(im).text((60, 250), "ACME Test Store INV-4417", fill="black", font=ImageFont.load_default(size=70))
        path = os.path.join(d, "receipt.jpg")
        im.save(path, "JPEG", **({"dpi": (density, density)} if density else {}))
        return path, d

    def convert(self, path, d, *flags):
        r = subprocess.run([sys.executable, "-W", "ignore", os.path.join(REPO, "tiff_to_pdf.py"), path, "--outdir", d, *flags],
                           capture_output=True, text=True, timeout=300)
        self.assertEqual(r.returncode, 0, r.stdout[-500:] + r.stderr[-300:])
        return pymupdf.open(os.path.join(d, "receipt.pdf"))

    def test_fit_a4_ignores_a_stored_jpeg_density(self):
        path, d = self.jpg(density=72)
        self.assertAlmostEqual(self.convert(path, d, "--no-text")[0].rect.width, 1000, delta=1,
                               msg="without the option a 72 dpi JPEG makes a 1000 pt page")
        self.assertAlmostEqual(self.convert(path, d, "--no-text", "--fit-a4")[0].rect.width, 842, delta=1)

    @unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
    def test_picture_is_exact_and_text_is_searchable(self):
        path, d = self.jpg()
        doc = self.convert(path, d, "--fit-a4")
        page = doc[0]
        self.assertEqual(len(page.get_images()), 1, "the picture is on the page")
        self.assertIn("INV-4417", page.get_text())
        self.assertTrue(page.search_for("INV-4417"))

    def test_the_route_is_registered_without_the_model(self):
        os.environ.setdefault("IDOX_JOBS_DIR", tempfile.mkdtemp())
        from api import idox_api
        route = idox_api.ROUTES["jpg_pdf_search"]
        self.assertFalse(route["model"])
        self.assertIn("--fit-a4", route["extra"])
        self.assertEqual(route["script"], "tiff_to_pdf.py")


@unittest.skipUnless(HAVE_TESSERACT, "needs Tesseract")
class ScannedLogoAndLook(unittest.TestCase):
    """KNOWN_ISSUES #4 and #12: black headings, a logo cut out of the scan, an image-sized Word file."""

    def page_png(self, pdf, dpi=125):
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "p.png")
        pymupdf.open(pdf)[0].get_pixmap(dpi=dpi).save(path)
        return path, d

    def detect(self, png, out_dir, cells=frozenset()):
        import ocr_layout
        _, lines = ocr_layout.ocr_lines(png)
        return ocr_layout.scan_graphics(png, lines, cells, out_dir, 0, 72.0 / 125), lines

    def test_finds_the_logo_box_on_the_real_scan(self):
        png, d = self.page_png(AGREEMENT_SCANNED)
        pics, _ = self.detect(png, d)
        self.assertEqual(len(pics), 1)
        x0, y0, x1, y1 = pics[0]["bbox_pt"]
        self.assertGreater(x0, 595 * 0.6, "top right of an A4 page")
        self.assertLess(y1, 842 * 0.15)
        self.assertTrue(pics[0]["float"])

    def test_finds_nothing_on_plain_pages(self):
        import scan_fixtures
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        for fn in (scan_fixtures.table_low, scan_fixtures.two_tables, scan_fixtures.wide_table,
                   scan_fixtures.no_table, scan_fixtures.table_high):
            pdf = os.path.join(d, fn.__name__ + ".pdf")
            fn(pdf)
            png, _ = self.page_png(pdf)
            self.assertEqual(self.detect(png, d)[0], [], fn.__name__)

    def test_a_shaded_header_band_or_a_boxed_paragraph_is_not_a_logo(self):
        import ocr_layout
        from PIL import Image, ImageDraw
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        cases = {"wide band": (60, 120, 900, 170),          # shaded table header across the page
                 "lower half": (700, 900, 900, 960),        # solid box low on the page
                 "huge box": (100, 100, 900, 600)}          # far too big for a logo
        for name, box in cases.items():
            im = Image.new("RGB", (1000, 1400), "white")
            ImageDraw.Draw(im).rectangle(box, fill=(120, 120, 120))
            png = os.path.join(d, name.replace(" ", "_") + ".png")
            im.save(png)
            self.assertEqual(ocr_layout.scan_graphics(png, [], frozenset(), d, 0, 0.5), [], name)

    def test_a_solid_box_with_a_paragraph_in_it_is_text_not_a_logo(self):
        import ocr_layout
        from PIL import Image, ImageDraw
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        im = Image.new("RGB", (1000, 1400), "white")
        ImageDraw.Draw(im).rectangle((650, 60, 900, 160), fill=(120, 120, 120))
        png = os.path.join(d, "box.png")
        im.save(png)
        lines = [{"x0": 660, "x1": 890, "top": 70 + 30 * i, "bottom": 90 + 30 * i, "text": "a line of real text"} for i in range(3)]
        self.assertEqual(ocr_layout.scan_graphics(png, lines, frozenset(), d, 0, 0.5), [])
        one = [{"x0": 700, "x1": 850, "top": 90, "bottom": 130, "text": "ACME"}]
        self.assertEqual(len(ocr_layout.scan_graphics(png, one, frozenset(), d, 0, 0.5)), 1, "the same box with one word is a logo")

    def test_the_logos_lettering_is_removed_from_the_text(self):
        import ocr_layout
        plan = {"sc": 0.5, "pages": [{"lines": [{"x0": 700, "x1": 850, "top": 90, "bottom": 130, "text": ". ACME"},
                                                 {"x0": 60, "x1": 400, "top": 100, "bottom": 140, "text": "Services Agreement"}]}]}
        pics = [[{"bbox_pt": (340, 40, 440, 70)}]]
        page = page_of([text("heading", "Services Agreement\nSample document\nACME"), text("paragraph", "Body.")])
        out = ocr_layout._strip_picture_lines([page], plan, pics)[0]
        self.assertEqual(out.blocks[0].text, "Services Agreement\nSample document")

    def test_scanned_headings_are_black_not_blue(self):
        import docx
        from to_docx import build_docx
        layout = {"width_pt": 595, "height_pt": 842, "left_pt": 72, "right_pt": 72, "top_pt": 72, "bottom_pt": 72, "body_pt": 11}
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "t.docx")
        build_docx([page_of([text("heading", "Title"), text("paragraph", "Body.")])], path, layout=layout)
        style = docx.Document(path).styles["Heading 1"]
        self.assertEqual(str(style.font.color.rgb), "000000")
        self.assertTrue(style.font.bold, "bold is kept for measured pages")

    def test_image_to_word_takes_its_size_font_and_table_from_the_image(self):
        import docx
        import ocr_layout
        from docx.shared import Pt
        d = tempfile.mkdtemp()
        self.addCleanup(shutil.rmtree, d, True)
        path = os.path.join(d, "t.docx")
        doc = docx.Document()
        doc.add_heading("Purchase list", 1)
        t = doc.add_table(rows=3, cols=3)
        for r, row in enumerate((("Item", "Qty", "Amount"), ("Paper A4", "10", "450.00"), ("Toner", "2", "3,200.00"))):
            for c, v in enumerate(row):
                t.cell(r, c).text = v
        doc.save(path)
        png = os.path.join(d, "img.png")
        from PIL import Image, ImageDraw
        im = Image.new("RGB", (1000, 600), "white")
        draw = ImageDraw.Draw(im)
        for y in (100, 200, 300, 400):
            draw.line((50, y, 950, y), fill="black", width=3)             # four ruling lines = three rows
        im.save(png)
        info = {"body_h": 30, "page_centre": 500}
        body = ocr_layout.polish_image_docx(path, page_of([text("heading", "Purchase list")]), info, {}, png)
        out = docx.Document(path)
        text_w = (out.sections[0].page_width - out.sections[0].left_margin - out.sections[0].right_margin) / 12700
        self.assertAlmostEqual(body, round(30 * (text_w / 1000) * ocr_layout.FONT_FROM_BOX * 2) / 2, delta=0.01)
        self.assertEqual(out.styles["Normal"].font.name, "Arial")
        self.assertAlmostEqual(sum(c.width for c in out.tables[0].columns) / 12700, text_w, delta=2)
        self.assertAlmostEqual(out.tables[0].rows[0].height / 12700, 100 * text_w / 1000, delta=1.5, msg="row height from the ruling lines")


class ModelAddress(unittest.TestCase):
    """Review item: the default model address can be set from the environment; unset, it is unchanged."""

    def default_in_subprocess(self, **env):
        e = {k: v for k, v in os.environ.items() if k != "IDOX_BASE_URL"}
        e.update(env)
        r = subprocess.run([sys.executable, "-W", "ignore", "-c", "import blocks; print(blocks.DEFAULT_BASE_URL)"],
                           capture_output=True, text=True, cwd=REPO, env=e, timeout=120)
        return r.stdout.strip().splitlines()[-1]

    def test_unset_keeps_the_old_default(self):
        self.assertEqual(self.default_in_subprocess(), "http://10.0.3.33:8080")

    def test_the_variable_wins(self):
        self.assertEqual(self.default_in_subprocess(IDOX_BASE_URL="http://127.0.0.1:8080"), "http://127.0.0.1:8080")


class ExcelNumbers(unittest.TestCase):
    """Review item: table amounts are real numbers that display exactly as printed; anything doubtful stays text."""

    def test_unambiguous_numbers_convert_and_round_trip(self):
        from to_xlsx import as_number, cell_display
        for text in ("450.00", "25,000.00", "1,250", "10", "-3.25", "0.5", "2026", "0", "1,000,000.50", "-1,250.00"):
            got = as_number(text)
            self.assertIsNotNone(got, text)
            self.assertEqual(cell_display(*got), text, "the number must display exactly as printed")

    def test_doubtful_values_stay_text(self):
        from to_xlsx import as_number
        for text in ("007", "+91 98765 43210", "$1,020.00", "12%", "1,25,000.00", "123456", "2026-10-05", "-0.00",
                     "1.2.3", "", "N/A", "=SUM(A1)", "-5%", " 450.00", "12/05/2026", "1e5", "0123.50",
                     "123456789012345678"):
            self.assertIsNone(as_number(text), repr(text))

    def build(self, **env):
        code = (
            "import sys, tempfile, os; sys.path.insert(0, %r)\n"
            "from blocks import *\nfrom to_xlsx import build_xlsx, verify_xlsx\nfrom openpyxl import load_workbook\n"
            "t = TableBlock(kind='table', has_header=True, n_data_rows=3, n_cols=3, header=['Item','Qty','Amount (INR)'],"
            " rows=[['Paper A4','10','450.00'],['Toner','2','3,200.00'],['Ref','007','2026-10-05']])\n"
            "pg = Page(analysis=LayoutAnalysis(n_blocks=2, table_column_counts=[3]), blocks=[TextBlock(kind='paragraph',"
            " text='Reference 2026', align='left', size='normal', bold=False), t])\n"
            "p = os.path.join(tempfile.mkdtemp(), 'x.xlsx'); build_xlsx([pg], p)\n"
            "ws = load_workbook(p).active\n"
            "cells = {c.coordinate: (type(c.value).__name__, c.value, c.number_format) for r in ws.iter_rows() for c in r if c.value is not None}\n"
            "import json; print(json.dumps({'cells': cells, 'coverage': verify_xlsx(p, 'Reference 2026 Paper A4 450.00 3,200.00 Amount Toner')[0]}))\n"
        ) % REPO
        e = {k: v for k, v in os.environ.items() if k != "IDOX_XLSX_NUMBERS"}
        e.update(env)
        r = subprocess.run([sys.executable, "-W", "ignore", "-c", code], capture_output=True, text=True, env=e,
                           timeout=120, cwd=REPO)
        self.assertEqual(r.returncode, 0, r.stderr[-500:])
        return json.loads(r.stdout.strip().splitlines()[-1])

    def test_workbook_has_numbers_where_safe_and_text_elsewhere(self):
        out = self.build()
        cells = out["cells"]
        self.assertEqual(cells["A1"][0], "str", "prose stays text")          # "Reference 2026"
        # row 1 is the prose, row 3 the header, rows 4 to 6 the data
        for coord, number, fmt in (("C4", 450, "0.00"), ("C5", 3200, "#,##0.00"), ("B4", 10, "0")):
            kind, value, shown = cells[coord]                                # a whole float is read back as int
            self.assertIn(kind, ("int", "float"), coord)
            self.assertEqual((value, shown), (number, fmt), coord)
        self.assertEqual(cells["A6"][0], "str")                              # a text cell in the table
        self.assertEqual(cells["B6"][:2], ["str", "007"], "leading zero stays text")
        self.assertEqual(cells["C6"][:2], ["str", "2026-10-05"], "a date stays text")
        self.assertEqual(cells["C3"][:2], ["str", "Amount (INR)"], "the header is always text")
        self.assertEqual(out["coverage"], 1.0, "reading the workbook back still finds every word")

    def test_the_old_behaviour_is_one_setting_away(self):
        cells = self.build(IDOX_XLSX_NUMBERS="0")["cells"]
        self.assertEqual(cells["C4"][:2], ["str", "450.00"])


# --------------------------------------------------------------------------- the model being unreachable
class ModelDown(unittest.TestCase):
    """A page the model cannot read must not look like success (review item: fail loudly)."""

    def test_exit_code_is_non_zero_when_every_page_fails(self):
        with tempfile.TemporaryDirectory() as d:
            r = subprocess.run([sys.executable, os.path.join(REPO, "test_pdf.py"), AGREEMENT, "--docx",
                                os.path.join(d, "x.docx"), "--base-url", "http://127.0.0.1:9"],
                               capture_output=True, text=True, timeout=300, cwd=d)
        self.assertIn("Tier 1 FAILED", r.stdout)
        self.assertIn("VERDICT: FAIL", r.stdout, "the log must say the run failed")
        self.assertNotIn("Traceback", r.stderr, "a crash is not the same as a deliberate failure")
        self.assertEqual(r.returncode, 1, "a run that read no page must exit 1, not 0")


if __name__ == "__main__":
    unittest.main()
