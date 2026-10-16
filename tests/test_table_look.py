"""A table keeps the look it had in the PDF: per-column alignment, the header's fill colour, column widths.

The look is measured from the PDF's own text layer and drawings (no model), then written by the Word, PowerPoint and
Excel writers. A synthetic PDF with a grey header and centred number columns stands in for a real one.
"""
import os
import tempfile
import unittest

import _load  # noqa: F401  (puts the repo on sys.path)
import pymupdf
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from test_no_model import page_of, table

import pdf_look
from to_docx import build_docx
from to_pptx import build_pptx
from to_xlsx import build_xlsx

HEADER = ["Department", "New Hires", "Pending"]
ROWS = [["Engineering", "18", "2"], ["Sales", "9", "0"]]
GRID = [HEADER] + ROWS


def synthetic_pdf():
    """One page: a 3-column table, first column left-aligned, the others centred, header shaded grey."""
    doc = pymupdf.open()
    pg = doc.new_page()
    x0, y0, widths, rh = 72, 100, [150, 100, 100], 24
    for r, row in enumerate(GRID):
        x = x0
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x, y0 + r * rh, x + widths[c], y0 + (r + 1) * rh)
            pg.draw_rect(rect, color=(0, 0, 0), fill=(0.85, 0.85, 0.85) if r == 0 else None, width=0.8)
            pg.insert_textbox(rect + (5, 6, -5, 0), value, fontsize=10, align=0 if c == 0 else 1)
            x += widths[c]
    return doc


class MeasuredLook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.doc = synthetic_pdf()
        page = page_of([table(HEADER, ROWS)])
        cls.styled = pdf_look.measure_tables([page], [cls.doc[0]], [True])[0]
        cls.block = cls.styled.blocks[0]

    def test_alignment_is_read_per_column(self):
        self.assertEqual(self.block.col_align, ["left", "center", "center"])
        self.assertEqual(self.block.header_align, ["left", "center", "center"])

    def test_header_fill_is_read(self):
        self.assertEqual(self.block.header_fill, "D9D9D9")

    def test_unshaded_table_has_no_fill(self):
        doc = pymupdf.open()
        pg = doc.new_page()
        for r, row in enumerate(GRID):
            for c, value in enumerate(row):
                rect = pymupdf.Rect(72 + c * 120, 100 + r * 24, 72 + (c + 1) * 120, 124 + r * 24)
                pg.draw_rect(rect, color=(0, 0, 0), width=0.8)
                pg.insert_textbox(rect + (5, 6, -5, 0), value, fontsize=10)
        block = pdf_look.measure_tables([page_of([table(HEADER, ROWS)])], [doc[0]], [True])[0].blocks[0]
        self.assertIsNone(block.header_fill)
        self.assertEqual(block.col_align, ["left", "left", "left"])

    def test_a_scanned_page_is_left_alone(self):
        page = page_of([table(HEADER, ROWS)])
        self.assertIs(pdf_look.measure_tables([page], [self.doc[0]], [False])[0], page)


class WritersApplyTheLook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.pages = pdf_look.measure_tables([page_of([table(HEADER, ROWS)])], [synthetic_pdf()[0]], [True])
        cls.dir = tempfile.TemporaryDirectory()

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def path(self, name):
        return os.path.join(self.dir.name, name)

    def test_word(self):
        build_docx(self.pages, self.path("t.docx"))
        t = Document(self.path("t.docx")).tables[0]
        self.assertEqual([c.paragraphs[0].alignment for c in t.rows[1].cells], [0, 1, 1])      # left, centre, centre
        self.assertIn('w:fill="D9D9D9"', t.rows[0].cells[0]._tc.xml)
        self.assertNotIn("w:fill=", t.rows[1].cells[0]._tc.xml)

    def test_powerpoint(self):
        build_pptx(self.pages, self.path("t.pptx"))
        shape = next(s for s in Presentation(self.path("t.pptx")).slides[0].shapes if s.has_table)
        table_ = shape.table
        self.assertEqual([c.text_frame.paragraphs[0].alignment for c in table_.rows[1].cells][1:], [2, 2])   # PP_ALIGN.CENTER == 2
        self.assertEqual(str(table_.cell(0, 0).fill.fore_color.rgb), "D9D9D9")
        widths = [c.width for c in table_.columns]
        self.assertGreater(widths[0], widths[1])                                                         # 150 pt vs 100 pt

    def test_excel(self):
        build_xlsx(self.pages, self.path("t.xlsx"))
        ws = load_workbook(self.path("t.xlsx")).active
        self.assertEqual(ws["B2"].alignment.horizontal, "center")
        self.assertEqual(ws["A2"].alignment.horizontal, "left")
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], "D9D9D9")

    def test_unmeasured_table_keeps_the_default_look(self):
        plain = [page_of([table(HEADER, ROWS)])]
        build_xlsx(plain, self.path("p.xlsx"))
        ws = load_workbook(self.path("p.xlsx")).active
        self.assertIsNone(ws["A1"].fill.fgColor.rgb if ws["A1"].fill.fill_type else None)


if __name__ == "__main__":
    unittest.main()


class ControlCharactersDoNotStopWord(unittest.TestCase):
    """A form feed or NUL in a PDF's text layer used to make the whole Word build fall back to default styles."""

    def test_word_builds_with_control_characters_in_text_and_cells(self):
        from test_no_model import text
        page = page_of([text("paragraph", "before\x0cafter\x00 end"), table(["A\x0b", "B"], [["1\x0c", "2"]])])
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "t.docx")
            build_docx([page], path)
            doc = Document(path)
            self.assertEqual(doc.paragraphs[0].text, "beforeafter end")
            self.assertEqual(doc.tables[0].rows[0].cells[0].text, "A")


class RunningFooterGoesLast(unittest.TestCase):
    """A PDF that stores its footer first must not put the footer first in Word (it pushes the page onto a blank one)."""

    def test_footer_listed_first_is_moved_to_the_end(self):
        from test_no_model import text
        doc = pymupdf.open()
        pg = doc.new_page()
        pg.insert_text((72, 815), "Sample Document footer", fontsize=8)           # written first, sits at the page bottom
        pg.insert_text((72, 80), "Report heading", fontsize=18)
        pg.insert_text((72, 120), "A body paragraph of ordinary text on the page.", fontsize=10)
        page = page_of([text("paragraph", "Sample Document footer"), text("heading", "Report heading"),
                        text("paragraph", "A body paragraph of ordinary text on the page.")])
        pdf_look.restyle_page(page, pg, False)
        self.assertEqual([b.text for b in page.blocks],
                         ["Report heading", "A body paragraph of ordinary text on the page.", "Sample Document footer"])
