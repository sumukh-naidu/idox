"""
show_docx.py -- print a .docx as readable text in the terminal.

A .docx is a zip archive of XML, so opening it in a text editor shows binary
noise. This reads the real structure back out and prints it, which is also the
honest way to check the conversion: it reports what is ACTUALLY IN THE FILE,
not what the converter meant to put there.

    .venv/bin/python show_docx.py output.docx
    .venv/bin/python show_docx.py output.docx --styles    # show Word styles too

Merged cells are reported as such -- a spanning divider row should appear as
one wide cell, not as the same text repeated across the row.
"""

import argparse

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

parser = argparse.ArgumentParser()
parser.add_argument("path")
parser.add_argument("--styles", action="store_true",
                    help="print the Word style name of each paragraph")
args = parser.parse_args()

doc = Document(args.path)


def iter_body(document):
    """Walk paragraphs and tables in the order they appear in the document.

    doc.paragraphs and doc.tables are separate lists, so neither shows where a
    table sits relative to the text around it. Walking the body XML preserves
    the real reading order.
    """
    body = document.element.body
    for child in body.iterchildren():
        if child.tag == f"{W}p":
            yield Paragraph(child, document)
        elif child.tag == f"{W}tbl":
            yield Table(child, document)


def show_table(table):
    # Column width from the grid definition, not from len(row.cells), which
    # counts a merged cell once per column it spans.
    grid = table._tbl.find(f"{W}tblGrid")
    ncols = len(grid.findall(f"{W}gridCol")) if grid is not None else 0
    rows = table._tbl.findall(f"{W}tr")
    print(f"  TABLE  {len(rows)} rows x {ncols} cols")

    for r, tr in enumerate(rows):
        tcs = tr.findall(f"{W}tc")
        cells = []
        merged = False
        for tc in tcs:
            span_el = tc.find(f".//{W}gridSpan")
            span = int(span_el.get(f"{W}val")) if span_el is not None else 1
            text = " ".join(t.text or "" for t in tc.iter(f"{W}t")).strip()
            if span > 1:
                merged = True
                cells.append(f"{text}  <-- merged across {span} columns")
            else:
                cells.append(text)
        tag = "  [spanning row]" if merged else ""
        print(f"    {r}: " + " | ".join(c[:28] for c in cells) + tag)


print(f"=== {args.path} ===\n")

counts = {"paragraphs": 0, "tables": 0}
for item in iter_body(doc):
    if isinstance(item, Table):
        counts["tables"] += 1
        show_table(item)
        print()
    else:
        # A picture lives inside a paragraph with no text, so checking
        # item.text alone reports an empty paragraph and hides the image.
        drawings = item._p.findall(f".//{W}drawing")
        if drawings:
            for d in drawings:
                # wp:extent, in the wordprocessingDrawing namespace -- not the
                # drawingml "main" one, which is where a:ext lives.
                ext = d.find(f".//{{http://schemas.openxmlformats.org/"
                             f"drawingml/2006/wordprocessingDrawing}}extent")
                if ext is not None:
                    w = int(ext.get("cx")) / 914400   # EMU per inch
                    h = int(ext.get("cy")) / 914400
                    print(f"  [IMAGE]  {w:.2f} x {h:.2f} in")
                else:
                    print("  [IMAGE]")
                counts["images"] = counts.get("images", 0) + 1
            continue

        text = item.text.strip()
        if not text:
            continue
        counts["paragraphs"] += 1
        style = item.style.name if item.style is not None else "?"
        prefix = f"[{style}] " if args.styles else ""

        if style.startswith("Heading") or style == "Title":
            print(f"\n## {prefix}{text}\n")
        elif style.startswith("List"):
            print(f"  - {prefix}{text}")
        else:
            print(f"  {prefix}{text}")

print(f"\n--- {counts['paragraphs']} paragraphs, {counts['tables']} table(s), "
      f"{counts.get('images', 0)} image(s) ---")
