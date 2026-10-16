"""Tests for how to_pptx.build_pptx fits content onto slides. None of them calls the model.
Run from the idox folder:  python -m unittest discover -s tests -p "test_pptx_slides.py" -v
"""
import os
import tempfile
import unittest

import _load  # noqa: F401  (puts the repo on sys.path)
from pptx import Presentation
from test_no_model import page_of, table, text

from to_pptx import build_pptx, verify_pptx


class Slides(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.dir.cleanup()

    def build(self, blocks, name="t"):
        path = os.path.join(self.dir.name, name + ".pptx")
        problems = build_pptx([page_of(blocks)], path)
        return Presentation(path), problems, path

    def test_title_is_full_width_centred_and_on_the_slide(self):
        prs, _, _ = self.build([text("heading", "Quarterly Report"), text("paragraph", "Body text")])
        title = prs.slides[0].shapes.title
        self.assertEqual((title.left, title.width), (0, prs.slide_width))
        self.assertGreaterEqual(title.top, 0)
        self.assertGreater(title.height, 0)
        self.assertLess(title.top + title.height, prs.slide_height // 3)

    @staticmethod
    def tables_of(prs):
        return [sh.table for s in prs.slides for sh in s.shapes if sh.has_table]

    def test_a_page_that_fits_is_exactly_one_slide(self):
        prs, problems, _ = self.build([text("heading", "Short"), text("paragraph", "A few words here.")])
        self.assertEqual(len(prs.slides), 1)
        self.assertEqual(problems, [])

    def test_a_dense_page_continues_on_cont_slides_instead_of_being_squeezed(self):
        long_p = " ".join(f"Sentence number {i} has several plain words in it." for i in range(120))
        prs, problems, path = self.build([text("heading", "Dense Report"), text("paragraph", long_p)])
        titles = [s.shapes.title.text for s in prs.slides]
        self.assertGreater(len(prs.slides), 1)
        self.assertEqual(titles[0], "Dense Report")
        self.assertTrue(all(t == "Dense Report (cont.)" for t in titles[1:]))
        self.assertEqual(problems, [], "nothing should be reported as overcrowded: it was split, not squeezed")
        coverage, missing = verify_pptx(path, long_p)
        self.assertEqual(missing, [])
        self.assertEqual(coverage, 1.0)

    def test_text_is_split_between_lines_before_it_is_split_inside_one(self):
        paras = [text("paragraph", " ".join(f"Para {p} word {i}." for i in range(60))) for p in range(8)]
        prs, _, path = self.build([text("heading", "Many Paragraphs")] + paras)
        self.assertGreater(len(prs.slides), 1)
        coverage, missing = verify_pptx(path, " ".join(b.text for b in paras))
        self.assertEqual(missing, [])

    def test_a_long_table_splits_by_rows_and_repeats_its_header(self):
        rows = [[f"item{i}", str(i), f"{i}.00"] for i in range(40)]
        prs, problems, _ = self.build([text("heading", "Big Table"), table(["Name", "Qty", "Price"], rows)])
        tables = self.tables_of(prs)
        self.assertGreater(len(tables), 1)
        self.assertTrue(all(t.cell(0, 0).text == "Name" for t in tables), "every part keeps the header row")
        self.assertEqual(sum(len(t.rows) - 1 for t in tables), 40, "every row is on exactly one slide")
        self.assertTrue(all(len(t.rows) - 1 >= 3 for t in tables), "no slide is left with one or two orphan rows")
        self.assertEqual(problems, [])

    def test_a_table_that_fits_a_fresh_slide_is_moved_whole_not_split(self):
        rows = [[f"r{i}", str(i)] for i in range(10)]
        intro = [text("paragraph", " ".join(f"Intro sentence {i} fills the slide." for i in range(20)))]
        prs, _, _ = self.build([text("heading", "Mixed")] + intro + [table(["A", "B"], rows)])
        tables = self.tables_of(prs)
        self.assertEqual(len(tables), 1, "a table that fits a fresh slide must not be cut in two")
        self.assertEqual(len(tables[0].rows) - 1, 10)

    def test_a_small_table_beside_text_stays_on_one_slide(self):
        prs, problems, _ = self.build([text("heading", "T"), text("paragraph", "intro"),
                                       table(["A", "B"], [["1", "2"]] * 5)])
        self.assertEqual(len(prs.slides), 1)
        self.assertEqual(problems, [])

    def test_two_pages_keep_their_own_slides(self):
        path = os.path.join(self.dir.name, "two.pptx")
        build_pptx([page_of([text("heading", "One"), text("paragraph", "a")]),
                    page_of([text("heading", "Two"), text("paragraph", "b")])], path)
        self.assertEqual([s.shapes.title.text for s in Presentation(path).slides], ["One", "Two"])


if __name__ == "__main__":
    unittest.main()
