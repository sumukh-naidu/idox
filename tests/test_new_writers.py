"""Tests for the CSV and XML writers and the new conversion scripts. None of them calls the model.

Run from the idox folder:  python -m unittest tests.test_new_writers -v
"""
import csv
import importlib
import os
import tempfile
import unittest
import xml.etree.ElementTree as ET

import _load  # noqa: F401  (puts the repo on sys.path)
from test_no_model import page_of, table, text

from to_csv import build_csv, verify_csv
from to_xml import build_xml, verify_xml


def rows_of(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.reader(fh))


class CsvWriter(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.dir.name, "t.csv")

    def tearDown(self):
        self.dir.cleanup()

    def test_header_and_rows(self):
        build_csv([page_of([table(["Item", "Price"], [["Pen", "1.20"], ["Ink", "3.00"]])])], self.path)
        self.assertEqual(rows_of(self.path), [["Item", "Price"], ["Pen", "1.20"], ["Ink", "3.00"]])

    def test_text_is_one_line_per_row_in_reading_order(self):
        build_csv([page_of([text("heading", "Report"), text("paragraph", "First part\nsecond  part"),
                            table(["A", "B"], [["1", "2"]]), text("paragraph", "Closing words")])], self.path)
        self.assertEqual(rows_of(self.path), [["Report"], ["First part second part"], [], ["A", "B"], ["1", "2"], [],
                                              ["Closing words"]])

    def test_list_is_one_row_per_item_without_bullets(self):
        build_csv([page_of([text("list", "- apples\n* pears\n1. plums")])], self.path)
        self.assertEqual(rows_of(self.path), [["apples"], ["pears"], ["plums"]])

    def test_text_only_page_is_a_valid_csv(self):
        problems = build_csv([page_of([text("heading", "Hello"), text("paragraph", "a demo")])], self.path)
        self.assertEqual(rows_of(self.path), [["Hello"], ["a demo"]])
        self.assertEqual(problems, [])

    def test_two_tables_have_a_blank_row_between(self):
        build_csv([page_of([table(["A"], [["1"]])]), page_of([table(["B"], [["2"]])])], self.path)
        self.assertEqual(rows_of(self.path), [["A"], ["1"], [], ["B"], ["2"]])

    def test_ragged_row_is_padded_and_reported(self):
        problems = build_csv([page_of([table(["A", "B", "C"], [["1", "2"]])])], self.path)
        self.assertEqual(rows_of(self.path)[1], ["1", "2", ""])
        self.assertTrue(any("padded" in p for p in problems))

    def test_empty_row_is_dropped_and_reported(self):
        problems = build_csv([page_of([table(["Item", "Value"], [["", ""], ["Pages", "1"], ["  ", ""]])])], self.path)
        self.assertEqual(rows_of(self.path), [["Item", "Value"], ["Pages", "1"]])
        self.assertEqual(sum("empty row dropped" in p for p in problems), 2)

    def test_nothing_to_write_writes_nothing(self):
        problems = build_csv([page_of([text("paragraph", "   ")])], self.path)
        self.assertFalse(os.path.exists(self.path))
        self.assertTrue(any("nothing to write" in p for p in problems))

    def test_commas_and_quotes_survive(self):
        build_csv([page_of([table(["Note"], [['He said "hi", twice']])])], self.path)
        self.assertEqual(rows_of(self.path)[1], ['He said "hi", twice'])

    def test_byte_order_mark_for_excel(self):
        build_csv([page_of([table(["A"], [["é"]])])], self.path)
        with open(self.path, "rb") as fh:
            self.assertEqual(fh.read(3), b"\xef\xbb\xbf")

    def test_verify_reports_missing_words(self):
        build_csv([page_of([text("paragraph", "Notebook on the desk")])], self.path)
        coverage, missing = verify_csv(self.path, "Notebook desk Zeppelin")
        self.assertEqual(missing, ["zeppelin"])
        self.assertAlmostEqual(coverage, 2 / 3)

    def test_verify_clean(self):
        build_csv([page_of([text("heading", "Notebook"), table(["Item"], [["Stapler"]])])], self.path)
        self.assertEqual(verify_csv(self.path, "Notebook Stapler")[0], 1.0)


class XlsxProse(unittest.TestCase):
    def test_prose_cells_wrap_and_tables_do_not(self):
        from openpyxl import load_workbook
        from to_xlsx import build_xlsx
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "t.xlsx")
            build_xlsx([page_of([text("paragraph", "Title\nA long paragraph " * 5), table(["A", "B"], [["1", "2"]])])], path)
            ws = load_workbook(path).active
            self.assertTrue(ws["A1"].alignment.wrap_text)
            self.assertFalse(ws["A3"].alignment.wrap_text)


class XmlWriter(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.TemporaryDirectory()
        self.path = os.path.join(self.dir.name, "t.xml")

    def tearDown(self):
        self.dir.cleanup()

    def build(self, *pages):
        build_xml(list(pages), self.path, source_name="in.pdf")
        return ET.parse(self.path).getroot()

    def test_structure(self):
        root = self.build(page_of([text("heading", "Title"), text("paragraph", "Body text"),
                                   table(["A", "B"], [["1", "2"]])]))
        self.assertEqual(root.tag, "document")
        self.assertEqual(root.get("source"), "in.pdf")
        self.assertEqual(root.get("pages"), "1")
        page = root.find("page")
        self.assertEqual([c.tag for c in page], ["heading", "paragraph", "table"])
        self.assertEqual([c.text for c in page.find("table/header")], ["A", "B"])
        self.assertEqual([c.text for c in page.find("table/row")], ["1", "2"])

    def test_pages_in_order(self):
        root = self.build(page_of([text("paragraph", "first")]), page_of([text("paragraph", "second")]))
        self.assertEqual([p.get("number") for p in root.findall("page")], ["1", "2"])
        self.assertEqual(root.findall("page")[1].find("paragraph").text, "second")

    def test_list_items_lose_their_bullets(self):
        root = self.build(page_of([text("list", "- apples\n* pears\n1. plums")]))
        self.assertEqual([i.text for i in root.iter("item")], ["apples", "pears", "plums"])

    def test_special_characters_are_escaped_and_control_codes_dropped(self):
        root = self.build(page_of([text("paragraph", "a < b & c > d \x00\x0b end")]))
        self.assertEqual(root.find("page/paragraph").text, "a < b & c > d  end")

    def test_empty_block_is_skipped_and_reported(self):
        problems = build_xml([page_of([text("paragraph", "   ")])], self.path)
        self.assertTrue(any("empty" in p for p in problems))
        self.assertIsNone(ET.parse(self.path).getroot().find("page/paragraph"))

    def test_verify_reports_missing_words(self):
        build_xml([page_of([text("paragraph", "Notebook on the desk")])], self.path)
        coverage, missing = verify_xml(self.path, "Notebook desk Zeppelin")
        self.assertEqual(missing, ["zeppelin"])
        self.assertAlmostEqual(coverage, 2 / 3)


class CsvPlan(unittest.TestCase):
    """apply_plan copies every row; it only adds a title and decides whether the first row is a header."""

    @classmethod
    def setUpClass(cls):
        import csv_plan
        cls.apply = staticmethod(csv_plan.apply_plan)

    def test_headerless_file_keeps_its_first_row(self):
        out = self.apply(["North", "120"], [["South", "95"]], {"has_header": False, "title": ""})
        self.assertEqual(out, [["North", "120"], ["South", "95"]])

    def test_header_file_keeps_header_first(self):
        out = self.apply(["Region", "Units"], [["North", "120"]], {"has_header": True, "title": "Sales"})
        self.assertEqual(out, [[], ["Sales"], [], ["Region", "Units"], ["North", "120"]])

    def test_title_can_be_left_out(self):
        out = self.apply(["A"], [["1"]], {"has_header": True, "title": "Made up"}, with_title=False)
        self.assertEqual(out, [["A"], ["1"]])


class ScriptsImport(unittest.TestCase):
    """Each new script has a __main__ guard, so importing it must do nothing (no argparse, no model call)."""

    def test_every_new_script_imports_cleanly(self):
        for name in ("image_to_ppt", "image_to_xml", "image_to_csv", "pdf_to_xml", "pdf_to_csv", "word_to_image",
                     "word_to_xlsx", "word_to_xml", "word_to_ppt", "word_to_csv"):
            with self.subTest(name):
                importlib.import_module(name)

    def test_every_new_script_has_its_folders(self):
        for name in ("image_to_ppt", "image_to_xml", "image_to_csv", "pdf_to_xml", "pdf_to_csv", "word_to_image",
                     "word_to_xlsx", "word_to_xml", "word_to_ppt", "word_to_csv"):
            for suffix in ("input", "output"):
                with self.subTest(name + suffix):
                    self.assertTrue(os.path.isdir(os.path.join(_load.REPO, suffix, f"{name}_{suffix}")))


if __name__ == "__main__":
    unittest.main()
