"""A table read from a PICTURE keeps the picture's look: column widths, per-column alignment, header fill.

Measured from pixels only (no OCR, no model). A synthetic picture with a navy header and mixed alignments stands in for
a real one; the two sample pictures in input/ are checked for their grids too.
"""
import os
import tempfile
import unittest

import _load  # noqa: F401  (puts the repo on sys.path)
from docx import Document
from openpyxl import load_workbook
from PIL import Image, ImageDraw
from test_no_model import page_of, table, text

import image_look
from to_docx import build_docx
from to_xlsx import build_xlsx

HEADER = ["Region", "Units", "Revenue"]
ROWS = [["North", "120", "4500"], ["South", "95", "3800"], ["West", "60", "2100"]]


def table_picture(path, fill=(23, 54, 93)):
    """800x400 picture: ruled 3-column table, navy header with white text, text left / centre / right by column."""
    im = Image.new("RGB", (800, 400), "white")
    d = ImageDraw.Draw(im)
    xs, ys = [50, 300, 500, 700], [100, 140, 180, 220, 260]
    if fill:
        d.rectangle([xs[0], ys[0], xs[-1], ys[1]], fill=fill)
    for y in ys:
        d.line([xs[0], y, xs[-1], y], fill="black", width=2)
    for x in xs:
        d.line([x, ys[0], x, ys[-1]], fill="black", width=2)

    def put(c, r, value, how):
        w = d.textlength(value)
        x0, x1 = xs[c], xs[c + 1]
        x = x0 + 8 if how == "l" else (x0 + x1) / 2 - w / 2 if how == "c" else x1 - 8 - w
        d.text((x, ys[r] + 14), value, fill="white" if (r == 0 and fill) else "black")
    for c, value in enumerate(HEADER):
        put(c, 0, value, "l" if c == 0 else "c")
    for r, row in enumerate(ROWS, start=1):
        for c, value in enumerate(row):
            put(c, r, value, "l" if c == 0 else "r" if c == 2 else "c")
    im.save(path)


class PictureLook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.TemporaryDirectory()
        cls.path = os.path.join(cls.dir.name, "t.png")
        table_picture(cls.path)
        cls.page = image_look.apply_image_look(page_of([text("heading", "Sales"), table(HEADER, ROWS)]), cls.path)
        cls.block = cls.page.blocks[1]

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def test_grid_is_found(self):
        with Image.open(self.path) as im:
            grid = image_look.find_grids(im)[0]
        self.assertEqual(len(grid["cols"]), 4)
        self.assertEqual(len(grid["rows"]), 5)

    def test_widths_alignment_and_fill_are_measured(self):
        self.assertEqual([round(w) for w in self.block.col_widths], [250, 200, 200])
        self.assertEqual(self.block.col_align, ["left", "center", "right"])
        self.assertEqual(self.block.header_align, ["left", "center", "center"])
        self.assertEqual(self.block.header_fill, "17365D")
        self.assertEqual(self.block.header_text, "FFFFFF")

    def test_the_text_block_is_untouched(self):
        self.assertEqual(self.page.blocks[0].text, "Sales")

    def test_a_header_without_fill_has_none(self):
        path = os.path.join(self.dir.name, "plain.png")
        table_picture(path, fill=None)
        block = image_look.apply_image_look(page_of([table(HEADER, ROWS)]), path).blocks[0]
        self.assertIsNone(block.header_fill)
        self.assertEqual(block.col_align, ["left", "center", "right"])

    def test_a_grid_with_the_wrong_column_count_is_left_alone(self):
        wrong = table(["Region", "Units"], [["North", "120"], ["South", "95"], ["West", "60"]])
        page = page_of([wrong])
        result = image_look.apply_image_look(page, self.path)
        self.assertNotIsInstance(result.blocks[0], image_look.LookTable)
        self.assertEqual(result.blocks[0], wrong)

    def test_a_picture_without_a_table_returns_the_same_page(self):
        page = page_of([text("paragraph", "no table")])
        self.assertIs(image_look.apply_image_look(page, self.path), page)

    def test_a_missing_file_returns_the_same_page(self):
        page = page_of([table(HEADER, ROWS)])
        self.assertIs(image_look.apply_image_look(page, "no_such.png"), page)


class WritersApplyThePictureLook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = tempfile.TemporaryDirectory()
        path = os.path.join(cls.dir.name, "t.png")
        table_picture(path)
        cls.pages = [image_look.apply_image_look(page_of([table(HEADER, ROWS)]), path)]

    @classmethod
    def tearDownClass(cls):
        cls.dir.cleanup()

    def test_word(self):
        out = os.path.join(self.dir.name, "t.docx")
        build_docx(self.pages, out)
        t = Document(out).tables[0]
        self.assertEqual([c.paragraphs[0].alignment for c in t.rows[1].cells], [0, 1, 2])      # left, centre, right
        self.assertIn('w:fill="17365D"', t.rows[0].cells[0]._tc.xml)
        self.assertIn("FFFFFF", t.rows[0].cells[0]._tc.xml)                                      # white header text

    def test_excel(self):
        out = os.path.join(self.dir.name, "t.xlsx")
        build_xlsx(self.pages, out)
        ws = load_workbook(out).active
        self.assertEqual(ws["C2"].alignment.horizontal, "right")
        self.assertEqual(ws["A1"].fill.fgColor.rgb[-6:], "17365D")
        self.assertEqual(ws["A1"].font.color.rgb[-6:], "FFFFFF")


class SamplePictures(unittest.TestCase):
    """The grids of two real pictures from input/ (their file names are part of the repo)."""

    def grid(self, name):
        path = os.path.join(_load.REPO, "input", "image_to_csv_input", name)
        with Image.open(path) as im:
            return image_look.find_grids(im)

    def test_two_column_table(self):
        grids = self.grid("fol.png")
        self.assertEqual((len(grids), len(grids[0]["cols"]) - 1, len(grids[0]["rows"]) - 1), (1, 2, 4))

    def test_dense_text_is_not_mistaken_for_tables(self):
        grids = self.grid("demo2.png")
        self.assertEqual((len(grids), len(grids[0]["cols"]) - 1), (1, 4))


if __name__ == "__main__":
    unittest.main()
