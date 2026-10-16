"""
make_word_samples.py -- write the sample .docx files the Word conversions are tested on.

*.docx is git-ignored in this project, so these are rebuilt rather than committed:
    .venv/Scripts/python python_scripts/make_word_samples.py
Writes sample_report.docx (heading, paragraph, bullet list, table) and price_list.docx (heading and a table)
into every word_to_*_input/ folder. Existing files are left alone.
"""

import paths
import os

from docx import Document

FOLDERS = [paths.inp(f"word_to_{x}_input") for x in ("image", "xlsx", "xml", "ppt", "csv")]


def sample_report(path):
    d = Document()
    d.add_heading("Quarterly Report", level=1)
    d.add_paragraph("This report summarises sales for the first quarter. Revenue grew in every region.")
    d.add_heading("Highlights", level=2)
    for line in ("Revenue up twelve percent", "Two new offices opened", "Customer churn fell below three percent"):
        d.add_paragraph(line, style="List Bullet")
    d.add_heading("Sales by region", level=2)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([("Region", "Units", "Revenue"), ("North", "120", "24,000.00"),
                             ("South", "95", "19,500.00"), ("West", "143", "31,250.00")]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    d.save(path)


def price_list(path):
    d = Document()
    d.add_heading("Price List", level=1)
    t = d.add_table(rows=5, cols=3)
    t.style = "Table Grid"
    for r, row in enumerate([("Item", "Code", "Price"), ("Notebook", "NB-01", "4.50"), ("Pen", "PN-17", "1.20"),
                             ("Stapler", "ST-03", "8.75"), ("Folder", "FD-22", "0.90")]):
        for c, value in enumerate(row):
            t.cell(r, c).text = value
    d.save(path)


if __name__ == "__main__":
    for folder in FOLDERS:
        os.makedirs(folder, exist_ok=True)
        for name, build in (("sample_report.docx", sample_report), ("price_list.docx", price_list)):
            path = os.path.join(folder, name)
            if not os.path.exists(path):
                build(path)
                print("wrote", path)
