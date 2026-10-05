#!/usr/bin/env python3
"""
benchmark.py -- latency and accuracy of every idox conversion on four fixed inputs,
so two machines can be compared like for like.

    .venv/bin/python benchmark.py run            run everything on THIS machine
    .venv/bin/python benchmark.py score DIR      score a folder of outputs (from any machine)
    .venv/bin/python benchmark.py report DIR     rebuild report.md from DIR

THE FOUR INPUTS (the same on both machines)
    pdf_to_word_input/text_only_1page.pdf     a dense one-page digital PDF, text only
    pdf_to_word_input/sample.pdf              one page: headings, a bullet list, a small table
    image_to_word_input/hi.png                two short lines on a dark background
    benchmark_input/sample_125dpi.png         sample.pdf rendered at 125 dpi

WHAT "run" DOES
    - restarts the llama.cpp server before EVERY model conversion, so its first run is cold;
    - runs each conversion 3 times and records every run, not just an average;
    - takes prefill and generation time and token counts from the SERVER'S OWN LOG, not
      from the scripts' printouts, and flags runs whose image was served from the prompt
      cache (their prefill is about 0 and understates a first read);
    - records the hardware, the server build and flags, the model files' hashes and the
      inputs' hashes in machine.json, so a difference between machines can be explained;
    - keeps every output file: <out>/<input>/<conversion>/run<N>/out/...

SERVER SETTINGS USED (identical on both machines): 512 image tokens (min and max), one slot,
CPU only, context 8192, threads left at the server's default, the same two model files.

HOW ACCURACY IS SCORED (the same code scores any machine's files; see score_run)
    Ground truth: the PDF's own text layer for the two PDFs and for the render of sample.pdf;
    for hi.png, the two lines read off the image by a person.
    For a conversion that produces text (Word, Excel, PowerPoint, TXT, PDF):
      recall     share of the ground truth's distinct words longer than 3 characters that appear
                 in the output (the project's own coverage rule)
      precision  share of the output's distinct words (same length rule) found in the ground truth
      sequence   word-order-aware similarity (difflib) of the two word lists, 0-100. This drops
                 when words are missing, invented, repeated or out of order. HEADLINE NUMBER.
      figures    "exact" only if every number token matches in both directions
      tables     for text with a table: share of ground-truth table cells (Word/Excel/PowerPoint)
                 or table rows (TXT/PDF) found in the output
    For picture conversions: PSNR against a fresh render (JPG/JPEG), pixel-exactness (TIFF, and the
    picture inside a PDF or Word file), and OCR recall of the ground truth's words.
"""

import argparse
import difflib
import hashlib
import json
import math
import os
import platform
import re
import shutil
import signal
import subprocess
import sys
import time
import urllib.request
from pathlib import Path

REPO = Path(__file__).resolve().parent
sys.path.insert(0, str(REPO))

HOME = Path.home()
LLAMA = HOME / ".local/opt/llama.cpp/llama-b11247/llama-server"
MODEL = REPO / "models_manual/Qwen3VL-2B-Instruct-Q4_K_M.gguf"
MMPROJ = REPO / "models_manual/mmproj-Qwen3VL-2B-Instruct-Q8_0.gguf"
PORT = 8090
BASE_URL = f"http://127.0.0.1:{PORT}"
SERVER_ARGS = ["--port", str(PORT), "--host", "127.0.0.1", "--no-webui", "--ctx-size", "8192",
               "--image-min-tokens", "512", "--image-max-tokens", "512", "--parallel", "1"]

INPUTS = {
    "text_only_1page.pdf": "pdf_to_word_input/text_only_1page.pdf",
    "sample.pdf": "pdf_to_word_input/sample.pdf",
    "hi.png": "image_to_word_input/hi.png",
    "sample_125dpi.png": "benchmark_input/sample_125dpi.png",
}
HI_PNG_TRUTH = "Hello, this is a demo!\nGenerated as a PNG image with text"

PDF_MODEL = ["pdf_docx", "pdf_xlsx", "pdf_pptx", "pdf_txt"]
PDF_PLAIN = ["pdf_jpg", "pdf_jpeg", "pdf_tiff"]
IMG_MODEL = ["img_docx", "img_pdf", "img_xlsx"]
IMG_PLAIN = ["img_docx_pic"]
# (route, conversion it consumes the first output of)
DERIVED = {"pdf": [("docx_pdf", "pdf_docx"), ("tiff_pdf", "pdf_tiff"), ("tiff_pdf_pic", "pdf_tiff")],
           "image": [("docx_pdf", "img_docx")]}
SOURCE_EXT = {"docx_pdf": (".docx",), "tiff_pdf": (".tiff", ".tif"), "tiff_pdf_pic": (".tiff", ".tif")}


def kind_of(name):
    return "pdf" if name.endswith(".pdf") else "image"


def conv_dir(route, source=None):
    return route if not source else f"{route}@{source}"


# ===========================================================================
# ground truth and text extraction
# ===========================================================================

def pdf_truth(pdf_path):
    import pymupdf
    doc = pymupdf.open(pdf_path)
    text = "\n".join(p.get_text() for p in doc)
    tables, bullets, heads = [], 0, 0
    for page in doc:
        try:
            for t in page.find_tables().tables:
                rows = []
                for row in t.rows:
                    cells = [" ".join(page.get_text("text", clip=pymupdf.Rect(c)).split()) if c else ""
                             for c in row.cells]
                    cells = [c for c in cells if c]
                    if cells:
                        rows.append(cells)
                if rows:
                    tables.append(rows)
        except Exception:
            pass
        sizes = {}
        lines = []
        for b in page.get_text("dict")["blocks"]:
            if b["type"] != 0:
                continue
            for ln in b["lines"]:
                t = "".join(s["text"] for s in ln["spans"]).strip()
                if t:
                    sz = max(s["size"] for s in ln["spans"])
                    lines.append((t, sz))
                    sizes[round(sz)] = sizes.get(round(sz), 0) + len(t)
        body = max(sizes, key=sizes.get) if sizes else 10
        bullets += sum(1 for t, _ in lines if t.startswith("•"))
        heads += sum(1 for t, sz in lines if sz >= body * 1.25 and len(t) < 90)
    return {"text": text, "tables": tables, "bullets": bullets, "headings": heads}


def truth_for(input_name):
    if input_name == "hi.png":
        return {"text": HI_PNG_TRUTH, "tables": [], "bullets": 0, "headings": None}
    src = "sample.pdf" if input_name == "sample_125dpi.png" else input_name
    return pdf_truth(REPO / INPUTS[src])


def read_output(path):
    """(text, cells, info) of one output file. cells = list of table cell strings, or None."""
    ext = path.suffix.lower()
    info = {"tables": None, "bullets": None, "headings": None}
    if ext == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        d = Document(str(path))
        W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        parts, cells, heads, buls, ntab = [], [], 0, 0, 0
        for child in d.element.body.iterchildren():
            if child.tag == f"{W}p":
                p = Paragraph(child, d)
                if p.text.strip():
                    parts.append(p.text)
                sn = p.style.name if p.style is not None else ""
                heads += sn.startswith("Heading") or sn == "Title"
                buls += sn.startswith("List")
            elif child.tag == f"{W}tbl":
                ntab += 1
                for row in Table(child, d).rows:
                    seen, rowtxt = set(), []
                    for c in row.cells:
                        if id(c._tc) in seen:
                            continue
                        seen.add(id(c._tc))
                        t = c.text.strip()
                        if t:
                            cells.append(t)
                            rowtxt.append(t)
                    if rowtxt:
                        parts.append(" ".join(rowtxt))
        info.update(tables=ntab, bullets=buls, headings=heads, images=len(d.inline_shapes))
        return "\n".join(parts), cells, info
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(str(path), data_only=True)
        parts, cells = [], []
        from to_xlsx import cell_display      # numbers are real numbers now: compare the text they display
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                vals = [cell_display(c.value, c.number_format).strip() for c in row
                        if c.value is not None and cell_display(c.value, c.number_format).strip()]
                if vals:
                    parts.append(" ".join(vals)); cells.extend(vals)
        return "\n".join(parts), cells, info
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        parts, cells, ntab = [], [], 0
        for s in prs.slides:
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    parts.append(sh.text_frame.text)
                if sh.has_table:
                    ntab += 1
                    for row in sh.table.rows:
                        vals = [c.text.strip() for c in row.cells if c.text.strip()]
                        cells.extend(vals)
                        if vals:
                            parts.append(" ".join(vals))
        info["tables"] = ntab
        return "\n".join(parts), cells, info
    if ext == ".txt":
        lines = [l for l in path.read_text("utf-8", "replace").splitlines()
                 if not re.match(r"^----- Page \d+ -----$", l.strip())]
        info["bullets"] = sum(1 for l in lines if l.lstrip().startswith("- "))
        return "\n".join(lines), None, info
    if ext == ".pdf":
        import pymupdf
        return "\n".join(p.get_text() for p in pymupdf.open(str(path))), None, info
    return "", None, info


# ===========================================================================
# metrics
# ===========================================================================

def norm(s):
    from blocks import normalize
    return normalize(s)


def squash(s):
    from blocks import squash as sq
    return sq(s)


def long_words(text):
    return {w for w in norm(text).split() if len(w) > 3 and any(c.isalnum() for c in w)}


def number_tokens(text):
    return {w.strip(".,;:()[]") for w in norm(text).split() if any(c.isdigit() for c in w)} - {""}


def text_metrics(truth, out):
    g, o = norm(truth), norm(out)
    gsq, osq = squash(truth), squash(out)
    gw, ow = long_words(truth), long_words(out)
    rec = sum(1 for w in gw if w in o or squash(w) in osq) / len(gw) if gw else 1.0
    pre = sum(1 for w in ow if w in g or squash(w) in gsq) / len(ow) if ow else 0.0
    seq = difflib.SequenceMatcher(None, g.split(), o.split(), autojunk=False).ratio()
    gn, on = number_tokens(truth), number_tokens(out)
    return {"recall": round(rec, 4), "precision": round(pre, 4), "sequence": round(seq, 4),
            "f1": round(2 * rec * pre / (rec + pre), 4) if rec + pre else 0.0,
            "figures_exact": gn == on, "figures_missing": sorted(gn - on)[:6], "figures_extra": sorted(on - gn)[:6],
            "words_truth": len(g.split()), "words_out": len(o.split())}


def table_match(tables, cells, text):
    if not tables:
        return None
    if cells is not None:                              # Word / Excel / PowerPoint: whole cells
        have = {norm(c) for c in cells}
        want = [norm(c) for t in tables for r in t for c in r]
        return round(sum(1 for c in want if c in have) / len(want), 4) if want else None
    flat = norm(text)                                  # TXT / PDF: whole rows
    rows = [norm(" ".join(r)) for t in tables for r in t]
    return round(sum(1 for r in rows if r in flat) / len(rows), 4) if rows else None


def psnr(a, b):
    from PIL import ImageChops, ImageStat
    a, b = a.convert("RGB"), b.convert("RGB")
    if a.size != b.size:
        return None
    st = ImageStat.Stat(ImageChops.difference(a, b))
    rms = math.sqrt(sum(st.sum2) / (3 * a.width * a.height))
    return 99.0 if rms == 0 else round(20 * math.log10(255 / rms), 2)


def render_pdf(pdf_path, dpi):
    import pymupdf
    from PIL import Image
    pix = pymupdf.open(pdf_path)[0].get_pixmap(dpi=dpi, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def ocr_recall(image_path, truth):
    import ocr
    if not ocr.have_tesseract():
        return None
    got = norm(ocr.ocr_image(str(image_path)))
    gsq = squash(got)
    gw = long_words(truth)
    return round(sum(1 for w in gw if w in got or squash(w) in gsq) / len(gw), 4) if gw else None


def score_run(bench, input_name, conv, run_dir):
    """Accuracy of one run, from the OUTPUT FILES ONLY."""
    from PIL import Image
    out_dir = run_dir / "out"
    files = sorted(p for p in out_dir.rglob("*") if p.is_file() and not p.name.startswith(".")) if out_dir.exists() else []
    s = {"input": input_name, "conv": conv, "run": run_dir.name, "files": [p.relative_to(out_dir).as_posix() for p in files]}
    if not files:
        s["note"] = "no output file"
        return s
    truth = truth_for(input_name)
    route = conv.split("@")[0]
    in_pdf = REPO / INPUTS["sample.pdf" if input_name == "sample_125dpi.png" else input_name] if kind_of(input_name) == "pdf" or input_name == "sample_125dpi.png" else None
    try:
        if route in ("pdf_jpg", "pdf_jpeg"):
            jpgs = [p for p in files if p.suffix.lower() in (".jpg", ".jpeg")]
            ref = render_pdf(in_pdf, 150)
            im = Image.open(jpgs[0])
            s.update(pages=len(jpgs), psnr_db=psnr(im, ref), size_ok=im.size == ref.size,
                     ocr_recall=ocr_recall(jpgs[0], truth["text"]))
        elif route == "pdf_tiff":
            t = Image.open(files[0]); ref = render_pdf(in_pdf, 300)
            s.update(frames=getattr(t, "n_frames", 1), pixel_exact=(t.convert("RGB").tobytes() == ref.tobytes() and t.size == ref.size),
                     dpi=[float(x) for x in t.info.get("dpi", (0, 0))])
        elif route == "img_docx_pic":
            from docx import Document
            d = Document(str(files[0]))
            blob = next(iter(d.part.related_parts[r].blob for r in d.part.rels if "image" in d.part.rels[r].reltype), None)
            src = (REPO / INPUTS[input_name]).read_bytes()
            s.update(picture_identical=(blob == src), images_in_docx=len(d.inline_shapes))
        elif route in ("tiff_pdf", "tiff_pdf_pic"):
            import pymupdf
            src_name = conv.split("@")[1]
            tiff = next((p for p in (bench / input_name / src_name / "run1" / "out").rglob("*") if p.suffix.lower() in (".tiff", ".tif")), None)
            pdf = pymupdf.open(str(files[0]))
            pg = pdf[0]; imgs = pg.get_images(full=True)
            got = Image.open(__import__("io").BytesIO(pdf.extract_image(imgs[0][0])["image"])).convert("RGB") if imgs else None
            exact = None
            if tiff and got is not None:
                fr = Image.open(tiff).convert("RGB"); exact = (fr.tobytes() == got.tobytes() and fr.size == got.size)
            text = "\n".join(p.get_text() for p in pdf)
            s.update(pages=pdf.page_count, picture_exact=exact, has_text_layer=bool(text.strip()))
            if text.strip():
                s["text"] = text_metrics(truth["text"], text)
                s["text"]["table_match"] = table_match(truth["tables"], None, text)
        else:
            f = next((p for p in files if p.suffix.lower() in (".docx", ".xlsx", ".pptx", ".txt", ".pdf")), None)
            text, cells, info = read_output(f)
            m = text_metrics(truth["text"], text)
            m["table_match"] = table_match(truth["tables"], cells, text)
            m["tables_expected"] = len(truth["tables"]); m["tables_found"] = info.get("tables")
            m["bullets_expected"] = truth["bullets"]; m["bullets_found"] = info.get("bullets")
            m["headings_expected"] = truth["headings"]; m["headings_found"] = info.get("headings")
            s["text"] = m
            if route == "docx_pdf":                      # chain integrity: the Word file's own words survive into the PDF
                src_docx = next((p for p in (bench / input_name / conv.split("@")[1] / "run1" / "out").rglob("*.docx")), None)
                if src_docx:
                    dtext, _, _ = read_output(src_docx)
                    s["chain_recall_vs_docx"] = text_metrics(dtext, text)["recall"]
    except Exception as exc:
        s["error"] = f"{type(exc).__name__}: {exc}"
    return s


# ===========================================================================
# running
# ===========================================================================

def sh(cmd):
    return subprocess.run(cmd, capture_output=True, text=True).stdout.strip()


def sh_all(cmd):
    """stdout and stderr together (llama-server prints its version on stderr)."""
    r = subprocess.run(cmd, capture_output=True, text=True)
    return (r.stdout + r.stderr).strip()


def pkg_temp():
    for z in Path("/sys/class/thermal").glob("thermal_zone*"):
        try:
            if (z / "type").read_text().strip() == "x86_pkg_temp":
                return int((z / "temp").read_text()) // 1000
        except OSError:
            pass
    return None


def throttle_total():
    tot = 0
    for p in Path("/sys/devices/system/cpu").glob("cpu*/thermal_throttle/core_throttle_count"):
        try:
            tot += int(p.read_text())
        except (OSError, ValueError):
            pass
    return tot


def sha256(path, n=None):
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def machine_info(bench):
    cpuinfo = Path("/proc/cpuinfo").read_text()
    mem = dict(l.split(":", 1) for l in Path("/proc/meminfo").read_text().splitlines() if ":" in l)
    import importlib.metadata as md
    return {
        "when": time.strftime("%Y-%m-%d %H:%M:%S %z"),
        "cpu_model": re.search(r"model name\s*:\s*(.+)", cpuinfo).group(1),
        "logical_cpus": os.cpu_count(),
        "lscpu": {k: v.strip() for k, v in (l.split(":", 1) for l in sh(["lscpu"]).splitlines() if ":" in l)
                  if k in ("Thread(s) per core", "Core(s) per socket", "CPU max MHz", "CPU min MHz", "L3 cache")},
        "ram_total_mb": int(mem["MemTotal"].split()[0]) // 1024,
        "ram_available_mb_at_start": int(mem["MemAvailable"].split()[0]) // 1024,
        "governor": Path("/sys/devices/system/cpu/cpu0/cpufreq/scaling_governor").read_text().strip(),
        "os": f"{platform.system()} {platform.release()}",
        "llama_server": sh_all([str(LLAMA), "--version"]).replace("\n", " | ")[:200],
        "server_flags": [str(LLAMA), "--model", MODEL.name, "--mmproj", MMPROJ.name] + SERVER_ARGS,
        "server_threads": "server default (the log reports the number it chose)",
        "model_files": {p.name: {"bytes": p.stat().st_size, "sha256": sha256(p)} for p in (MODEL, MMPROJ)},
        "inputs": {n: {"path": p, "bytes": (REPO / p).stat().st_size, "sha256": sha256(REPO / p)} for n, p in INPUTS.items()},
        "python": sys.version.split()[0],
        "libs": {n: md.version(n) for n in ("pymupdf", "python-docx", "openpyxl", "python-pptx", "pillow", "pytesseract", "pydantic", "requests")},
        "tesseract": sh(["tesseract", "--version"]).splitlines()[0] if shutil.which("tesseract") else None,
        "libreoffice": sh(["soffice", "--version"]) if shutil.which("soffice") else None,
        "load_average_at_start": os.getloadavg(),
        "pkg_temp_c_at_start": pkg_temp(),
        "throttle_count_at_start": throttle_total(),
    }


def stop_server():
    pids = [int(x) for x in sh(["pgrep", "-x", "llama-server"]).split()]
    for pid in pids:
        try:
            os.kill(pid, signal.SIGTERM)
        except ProcessLookupError:
            pass
    t = time.time()
    while time.time() - t < 20 and sh(["pgrep", "-x", "llama-server"]):
        time.sleep(0.3)
    for pid in [int(x) for x in sh(["pgrep", "-x", "llama-server"]).split()]:
        os.kill(pid, signal.SIGKILL)


def start_server(logpath):
    stop_server()
    with open(logpath, "ab") as lg:
        lg.write(f"\n===== server (re)start {time.strftime('%H:%M:%S')} =====\n".encode())
        subprocess.Popen([str(LLAMA), "--model", str(MODEL), "--mmproj", str(MMPROJ)] + SERVER_ARGS,
                         stdout=lg, stderr=subprocess.STDOUT, stdin=subprocess.DEVNULL, start_new_session=True)
    t = time.time()
    while time.time() - t < 180:
        try:
            if urllib.request.urlopen(BASE_URL + "/health", timeout=2).status == 200:
                time.sleep(1.0)
                return True
        except Exception:
            time.sleep(0.5)
    raise RuntimeError("the model server did not become healthy")


RX_PRE = re.compile(r"prompt eval time =\s*([\d.]+) ms /\s*(\d+) tokens")
RX_GEN = re.compile(r"\|\s+eval time =\s*([\d.]+) ms /\s*(\d+) tokens")


def server_timing(logpath, start_offset):
    with open(logpath, "rb") as fh:
        fh.seek(start_offset)
        text = fh.read().decode("utf-8", "replace")
    pre = RX_PRE.findall(text); gen = RX_GEN.findall(text)
    if not pre:
        return None
    pre_s = sum(float(a) for a, _ in pre) / 1000; pre_n = sum(int(b) for _, b in pre)
    gen_s = sum(float(a) for a, _ in gen) / 1000; gen_n = sum(int(b) for _, b in gen)
    return {"requests": len(pre), "prefill_s": round(pre_s, 2), "prefill_tokens_evaluated": pre_n,
            "gen_s": round(gen_s, 2), "gen_tokens": gen_n, "tok_per_s": round(gen_n / gen_s, 2) if gen_s else None,
            "image_from_cache": pre_n < 100}


def run_one(bench, input_name, route, run_no, cold, source_file, model, logpath):
    import idox_app
    idox_app.CONFIG["backend"] = "local"
    cdir = conv_dir(route, source_file.parent.parent.parent.name if source_file else None)
    run_dir = bench / input_name / cdir / f"run{run_no}"
    if run_dir.exists():
        shutil.rmtree(run_dir)
    (run_dir / "in").mkdir(parents=True); (run_dir / "out").mkdir()
    src = source_file or (REPO / INPUTS[input_name])
    stem = idox_app.safe_stem(src.name)
    ext = src.suffix.lower().lstrip(".")
    shutil.copy2(src, run_dir / "in" / f"{stem}.{ext}")
    job = {"route": route, "stem": stem, "ext": ext, "options": {}, "base_url": BASE_URL}
    cmd = idox_app.build_command(job)
    load1 = os.getloadavg()[0]
    off = os.path.getsize(logpath) if model else 0
    t0 = time.perf_counter()
    with open(run_dir / "log.txt", "wb") as lg:
        try:
            rc = subprocess.run(cmd, cwd=run_dir, env=idox_app.job_env(model), stdout=lg, stderr=subprocess.STDOUT,
                                timeout=1500).returncode
        except subprocess.TimeoutExpired:
            rc = -9
    wall = time.perf_counter() - t0
    log = (run_dir / "log.txt").read_text("utf-8", "replace")
    outs = [p for p in (run_dir / "out").rglob("*") if p.is_file() and not p.name.startswith(".")]
    problems = [msg for rx, msg in idox_app.PROBLEMS if rx.search(log)]
    rec = {"input": input_name, "conv": cdir, "route": route, "model": model, "run": run_no, "cold_server": cold,
           "wall_s": round(wall, 2), "exit_code": rc,
           "state": "failed" if not outs else ("check" if problems or rc != 0 else "ok"),
           "outputs": len(outs), "problems": problems, "hit_token_cap": "generation limit" in log,
           "load1_before": round(load1, 2), "pkg_temp_c": pkg_temp(), "throttle_total": throttle_total(),
           "cmd": [Path(c).name if c.startswith("/") else c for c in cmd][1:]}
    if model:
        rec["server"] = server_timing(logpath, off)
    return rec


def first_output(bench, input_name, source_route, exts):
    d = bench / input_name / source_route / "run1" / "out"
    return next((p for p in sorted(d.rglob("*")) if p.is_file() and p.suffix.lower() in exts), None)


def cmd_run(a):
    bench = Path(a.out).resolve(); bench.mkdir(parents=True, exist_ok=True)
    logpath = bench / "server.log"
    runs_path = bench / "runs.jsonl"
    (bench / "machine.json").write_text(json.dumps(machine_info(bench), indent=1))
    done = set()
    if a.resume and runs_path.exists():
        for l in runs_path.read_text().splitlines():
            r = json.loads(l); done.add((r["input"], r["conv"], r["run"]))
    plan = []                                      # (input, route, source_route, model)
    for name in INPUTS:
        k = kind_of(name)
        for r in (PDF_MODEL if k == "pdf" else IMG_MODEL):
            plan.append((name, r, None, True))
        for r in (PDF_PLAIN if k == "pdf" else IMG_PLAIN):
            plan.append((name, r, None, False))
    for name in INPUTS:
        for r, src in DERIVED[kind_of(name)]:
            plan.append((name, r, src, False))
    if a.only:
        plan = [p for p in plan if re.search(a.only, f"{p[0]} {conv_dir(p[1], p[2])}")]
    print(f"{len(plan)} conversions x {a.runs} runs; output in {bench}", flush=True)
    t_all = time.time()
    for i, (name, route, srcroute, model) in enumerate(plan, 1):
        src_file = None
        if srcroute:
            src_file = first_output(bench, name, srcroute, SOURCE_EXT[route])
            if src_file is None:
                print(f"[{i}/{len(plan)}] {name} {route}@{srcroute}: no source output yet, skipped", flush=True)
                continue
        cname = conv_dir(route, srcroute)
        if model:
            if all((name, cname, r) in done for r in range(1, a.runs + 1)):
                continue
            start_server(logpath)
        for r in range(1, a.runs + 1):
            if (name, cname, r) in done:
                continue
            rec = run_one(bench, name, route, r, cold=(model and r == 1), source_file=src_file, model=model, logpath=logpath)
            with open(runs_path, "a") as fh:
                fh.write(json.dumps(rec) + "\n")
            sv = rec.get("server") or {}
            print(f"[{i}/{len(plan)}] {name:20s} {cname:22s} run{r} {rec['state']:6s} wall {rec['wall_s']:7.1f}s"
                  + (f" | prefill {sv.get('prefill_s')}s ({sv.get('prefill_tokens_evaluated')} tok{', cached' if sv.get('image_from_cache') else ''}) gen {sv.get('gen_s')}s/{sv.get('gen_tokens')} tok = {sv.get('tok_per_s')} tok/s" if sv else "")
                  + (" HIT TOKEN CAP" if rec["hit_token_cap"] else "") + f" | {pkg_temp()}C", flush=True)
    print(f"all runs finished in {(time.time() - t_all) / 60:.1f} min; scoring...", flush=True)
    cmd_score(argparse.Namespace(dir=str(bench)))


def cmd_score(a):
    bench = Path(a.dir).resolve()
    rows = []
    for ip in sorted(p for p in bench.iterdir() if p.is_dir() and p.name in INPUTS):
        for cp in sorted(p for p in ip.iterdir() if p.is_dir()):
            for rp in sorted(p for p in cp.iterdir() if p.is_dir() and p.name.startswith("run")):
                rows.append(score_run(bench, ip.name, cp.name, rp))
    with open(bench / "scores.jsonl", "w") as fh:
        for r in rows:
            fh.write(json.dumps(r, default=str) + "\n")
    print(f"scored {len(rows)} runs -> {bench / 'scores.jsonl'}", flush=True)
    cmd_report(argparse.Namespace(dir=str(bench)))


def mean(xs):
    xs = [x for x in xs if x is not None]
    return sum(xs) / len(xs) if xs else None


def f(x, d=1, pct=False):
    if x is None:
        return "-"
    return f"{x * 100:.{d}f}%" if pct else f"{x:.{d}f}"


def cmd_report(a):
    bench = Path(a.dir).resolve()
    runs = [json.loads(l) for l in (bench / "runs.jsonl").read_text().splitlines()] if (bench / "runs.jsonl").exists() else []
    scores = [json.loads(l) for l in (bench / "scores.jsonl").read_text().splitlines()] if (bench / "scores.jsonl").exists() else []
    machine = json.loads((bench / "machine.json").read_text()) if (bench / "machine.json").exists() else {}
    L = ["# idox benchmark report", ""]
    if machine:
        L += ["## Machine and settings", "",
              f"- CPU: {machine['cpu_model']} ({machine['logical_cpus']} logical CPUs; {machine['lscpu'].get('Core(s) per socket')} cores x {machine['lscpu'].get('Thread(s) per core')} threads; max {machine['lscpu'].get('CPU max MHz')} MHz), governor `{machine['governor']}`",
              f"- RAM: {machine['ram_total_mb']} MB total, {machine['ram_available_mb_at_start']} MB available at start",
              f"- OS: {machine['os']}; Python {machine['python']}",
              f"- Server: {machine['llama_server']}", f"- Server flags: `{' '.join(machine['server_flags'])}`",
              f"- Model files: " + "; ".join(f"{k} {v['bytes']} bytes sha256 {v['sha256'][:16]}" for k, v in machine['model_files'].items()),
              f"- Inputs: " + "; ".join(f"{k} sha256 {v['sha256'][:16]}" for k, v in machine['inputs'].items()),
              f"- Tesseract: {machine['tesseract']}; LibreOffice: {machine['libreoffice']}",
              f"- At start: load {machine['load_average_at_start'][0]:.2f}, package {machine['pkg_temp_c_at_start']} C, throttle events {machine['throttle_count_at_start']}", ""]
    if runs:
        temps = [r["pkg_temp_c"] for r in runs if r.get("pkg_temp_c") is not None]
        thr = [r["throttle_total"] for r in runs]
        loads = [r["load1_before"] for r in runs if r["model"]]
        L += ["## During the run", "",
              f"- {len(runs)} runs ({sum(1 for r in runs if r['model'])} with the model). Package temperature after each run: {min(temps)}-{max(temps)} C. Thermal throttle events: {thr[0]} after the first run, {thr[-1]} after the last.",
              f"- Load average (1 min) just before model runs: {min(loads):.2f}-{max(loads):.2f}. This includes the run that just finished and the server starting up, so it does not show what other programs were doing.",
              f"- Runs that hit the 4,000-token output cap: {sum(1 for r in runs if r['hit_token_cap'])}. Runs that produced no file: {sum(1 for r in runs if r['state'] == 'failed')} (the three `hi.png` to Excel runs are expected: that image has no table).", ""]
    by = {}
    for r in runs:
        by.setdefault((r["input"], r["conv"]), []).append(r)
    sc = {}
    for s in scores:
        sc.setdefault((s["input"], s["conv"]), []).append(s)
    L += ["## Latency, every run", "", "Wall = the whole script, seconds. Prefill / gen from the server's log (`cached` = image served from the prompt cache, so prefill is not a real first read).", "",
          "| input | conversion | model | run | state | wall s | prefill s | prefill tok | gen s | gen tok | tok/s | notes |", "|---|---|---|---|---|---|---|---|---|---|---|---|"]
    for (inp, conv), rs in by.items():
        for r in sorted(rs, key=lambda x: x["run"]):
            sv = r.get("server") or {}
            note = ("cold server; " if r["cold_server"] else "") + ("image cached; " if sv.get("image_from_cache") else "") + ("HIT 4000-TOKEN CAP; " if r["hit_token_cap"] else "") + "; ".join(r["problems"])
            L.append(f"| {inp} | {conv} | {'yes' if r['model'] else 'no'} | {r['run']} | {r['state']} | {r['wall_s']:.1f} | {f(sv.get('prefill_s'))} | {sv.get('prefill_tokens_evaluated', '-')} | {f(sv.get('gen_s'))} | {sv.get('gen_tokens', '-')} | {f(sv.get('tok_per_s'))} | {note} |")
    L += ["", "## Latency summary (per conversion)", "", "Cold = run 1 right after a server restart. Warm = mean of runs 2 and 3 (usually the image is cached).", "",
          "| input | conversion | model | cold wall s | warm wall s | cold prefill s | gen s (mean) | gen tok (mean) | tok/s (mean) | runs ok |", "|---|---|---|---|---|---|---|---|---|---|"]
    for (inp, conv), rs in by.items():
        rs = sorted(rs, key=lambda x: x["run"]); c = rs[0]; w = rs[1:]
        sv = [r.get("server") or {} for r in rs]
        L.append(f"| {inp} | {conv} | {'yes' if c['model'] else 'no'} | {c['wall_s']:.1f} | {f(mean([r['wall_s'] for r in w]))} | {f(sv[0].get('prefill_s'))} | {f(mean([s.get('gen_s') for s in sv]))} | {f(mean([s.get('gen_tokens') for s in sv]), 0)} | {f(mean([s.get('tok_per_s') for s in sv]))} | {sum(1 for r in rs if r['state'] != 'failed')}/{len(rs)} |")
    L += ["", "## Accuracy (per conversion, mean over runs; min in brackets for the headline)", "",
          "Headline = word-sequence similarity vs the ground truth (see benchmark.py). Recall/precision use words longer than 3 characters.", "",
          "| input | conversion | sequence % (min) | recall | precision | figures exact | table match | bullets found/expected | other |", "|---|---|---|---|---|---|---|---|---|"]
    for (inp, conv), ss in sc.items():
        tx = [s["text"] for s in ss if "text" in s]
        if tx:
            seqs = [t["sequence"] for t in tx]
            fe = sum(1 for t in tx if t["figures_exact"])
            tm = [t.get("table_match") for t in tx]
            tm = [x for x in tm if x is not None] or [None]
            b = tx[0]
            other = []
            ex = [s for s in ss if "chain_recall_vs_docx" in s]
            if ex: other.append(f"Word words kept in PDF {f(mean([s['chain_recall_vs_docx'] for s in ex]), 0, True)}")
            pe = [s.get('picture_exact') for s in ss if 'picture_exact' in s]
            if pe: other.append(f"picture exact {sum(1 for x in pe if x)}/{len(pe)}")
            if len(tx) < len(ss): other.append(f"{len(ss) - len(tx)} run(s) had no usable output")
            L.append(f"| {inp} | {conv} | {mean(seqs) * 100:.1f} ({min(seqs) * 100:.1f}) | {f(mean([t['recall'] for t in tx]), 1, True)} | {f(mean([t['precision'] for t in tx]), 1, True)} | {fe}/{len(tx)} | {f(mean(tm), 0, True) if mean(tm) is not None else 'n/a'} | {b.get('bullets_found') if b.get('bullets_found') is not None else '-'}/{b.get('bullets_expected', '-')} | {'; '.join(other)} |")
        else:
            bits = []
            for s in ss[:1]:
                for k in ("pages", "frames", "psnr_db", "size_ok", "ocr_recall", "pixel_exact", "picture_identical", "note", "error"):
                    if k in s: bits.append(f"{k} {s[k]}")
            L.append(f"| {inp} | {conv} | - | - | - | - | - | - | {'; '.join(bits)} |")
    L += ["", "## Accuracy, every run", "",
          "Run 1 follows a fresh server restart; runs 2 and 3 usually reuse the cached image. The model can give a different answer for the same input, so look at each run, not only the mean.", "",
          "| input | conversion | run | sequence % | recall | precision | figures exact | words out / truth | notes |", "|---|---|---|---|---|---|---|---|---|"]
    for s_ in scores:
        t = s_.get("text")
        if t:
            L.append(f"| {s_['input']} | {s_['conv']} | {s_['run']} | {t['sequence'] * 100:.1f} | {t['recall'] * 100:.1f}% | {t['precision'] * 100:.1f}% | {'yes' if t['figures_exact'] else 'NO ' + str(t['figures_missing'] + t['figures_extra'])} | {t['words_out']} / {t['words_truth']} | |")
    (bench / "report.md").write_text("\n".join(L) + "\n")
    # one flat row per run, for putting two machines side by side in a spreadsheet
    import csv
    sidx = {(x["input"], x["conv"], x["run"]): x for x in scores}
    with open(bench / "results.csv", "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["input", "conversion", "model", "run", "cold_server", "state", "wall_s", "prefill_s", "prefill_tokens_evaluated",
                    "image_from_cache", "gen_s", "gen_tokens", "tok_per_s", "hit_token_cap", "sequence_pct", "recall_pct", "precision_pct",
                    "figures_exact", "table_match_pct", "psnr_db", "ocr_recall_pct", "pixel_exact", "picture_exact_or_identical"])
        for r in runs:
            sv = r.get("server") or {}
            x = sidx.get((r["input"], r["conv"], f"run{r['run']}"), {})
            t = x.get("text") or {}
            pe = x.get("pixel_exact", x.get("picture_exact", x.get("picture_identical")))
            w.writerow([r["input"], r["conv"], r["model"], r["run"], r["cold_server"], r["state"], r["wall_s"], sv.get("prefill_s"),
                        sv.get("prefill_tokens_evaluated"), sv.get("image_from_cache"), sv.get("gen_s"), sv.get("gen_tokens"), sv.get("tok_per_s"),
                        r["hit_token_cap"], round(t["sequence"] * 100, 1) if t else None, round(t["recall"] * 100, 1) if t else None,
                        round(t["precision"] * 100, 1) if t else None, t.get("figures_exact") if t else None,
                        round(t["table_match"] * 100) if t.get("table_match") is not None else None, x.get("psnr_db"),
                        round(x["ocr_recall"] * 100) if x.get("ocr_recall") is not None else None, pe, ])
    print(f"report -> {bench / 'report.md'}", flush=True)


def main():
    global LLAMA, MODEL, MMPROJ, PORT, BASE_URL, SERVER_ARGS
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    r = sub.add_parser("run"); r.add_argument("--out", default=str(REPO / "benchmark_output")); r.add_argument("--runs", type=int, default=3)
    r.add_argument("--llama-server", help=f"path to llama-server (default {LLAMA})")
    r.add_argument("--model", help=f"path to the language model .gguf (default {MODEL})")
    r.add_argument("--mmproj", help=f"path to the vision projector .gguf (default {MMPROJ})")
    r.add_argument("--port", type=int, help=f"server port (default {PORT})")
    r.add_argument("--only", help="regex over '<input> <conversion>' to run a subset"); r.add_argument("--resume", action="store_true")
    s = sub.add_parser("score"); s.add_argument("dir")
    p = sub.add_parser("report"); p.add_argument("dir")
    a = ap.parse_args()
    if a.cmd == "run":
        LLAMA = Path(a.llama_server) if a.llama_server else LLAMA
        MODEL = Path(a.model) if a.model else MODEL
        MMPROJ = Path(a.mmproj) if a.mmproj else MMPROJ
        if a.port:
            PORT = a.port; BASE_URL = f"http://127.0.0.1:{PORT}"
            SERVER_ARGS = ["--port", str(PORT)] + SERVER_ARGS[2:]
    {"run": cmd_run, "score": cmd_score, "report": cmd_report}[a.cmd](a)


if __name__ == "__main__":
    main()
